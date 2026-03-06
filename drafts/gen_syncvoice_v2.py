#!/usr/bin/env python3
"""SyncVoice — Frontend + Backend decomposition (no status column, accurate to repo)."""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ─────────────────────────────────────────────────────────────────
# ДАННЫЕ: FRONTEND
# ─────────────────────────────────────────────────────────────────
# (phase, id, module, задачи, часы, стоимость, примечание)
FRONTEND = [
    # ── Phase 1: MVP UI ─────────────────────────────────────────
    ("Phase 1", "1.1", "Проект (Vite + TS)",
     "Инициализация React-проекта на Vite + TypeScript. Подключение @livekit/components-react, @livekit/client, настройка vite.config.ts, tsconfig, package.json, глобальные стили (global.css).",
     8, 20_000, "Vite build → static, Dockerfile для prod"),

    ("Phase 1", "1.2", "App.tsx — основной роутинг",
     "Главный компонент приложения: состояние подключения (token/connected), fetch токена с бэкенда (/token?room&identity&lang), интеграция LiveKitRoom, обработка onConnected / onDisconnected / onError, управление логами (addLog).",
     12, 30_000, ""),

    ("Phase 1", "1.3", "Controls",
     "Форма входа в комнату: поля room name, identity, выбор языка (en/ru/hi/zh), кнопки Join / Leave. Disabled-состояния при подключении. Валидация: блокировка Join без выбора языка.",
     10, 25_000, ""),

    ("Phase 1", "1.4", "AudioHandler",
     "Клиентская обработка аудио в LiveKit-комнате: подписка на треки, управление AudioContext. RoomAudioRenderer для воспроизведения TTS-треков от агента. Фильтрация sv_tts_* треков (не вешать на динамик дважды).",
     16, 40_000, ""),

    ("Phase 1", "1.5", "VideoGrid + VideoTile",
     "Сетка видео-тайлов участников: getOrCreateTile(), attachVideo(), removeTile(), setSpeaking(), toggleCamera(). Подписка на TrackSubscribed / TrackUnsubscribed / ParticipantConnected / ParticipantDisconnected. Включение камеры при подключении, LocalTrackPublished handler.",
     24, 60_000, "VideoGrid + VideoTile — CSS сделан, JS tiles не полностью"),

    ("Phase 1", "1.6", "ParticipantList",
     "Список участников в комнате: имя, язык (из metadata), иконка микрофона/камеры, индикатор «говорит».",
     8, 20_000, ""),

    ("Phase 1", "1.7", "EventLog",
     "Панель событий: лог последних 50 сообщений с timestamp, отображение ошибок / статусов подключения / событий перевода.",
     6, 15_000, ""),

    ("Phase 1", "1.8", "AudioDebugPanel",
     "Debug-панель в dev-режиме: уровень RMS входящего аудио, статус VAD на клиенте (если нужен), сетевые метрики LiveKit.",
     8, 20_000, ""),

    # ── Phase 2: Продукт ──────────────────────────────────────────
    ("Phase 2", "2.1", "Auth — UI",
     "Страницы логина и регистрации (JWT auth). Форма для гостевого входа по ссылке. React Router: защита роутов по ролям (Админ, Модератор, Участник, Гость). Persist токена в localStorage.",
     24, 60_000, "Зависит от Backend Auth API"),

    ("Phase 2", "2.2", "Смена языка в комнате",
     "UI переключения языка во время конференции. Отправка нового metadata в LiveKit (participant.setMetadata). Перезапрос TTS-треков у агента.",
     12, 30_000, "Требует поддержки на бэкенде (agent перезапускает pipeline)"),

    ("Phase 2", "2.3", "Screen Sharing",
     "Кнопка демонстрации экрана, интеграция с LiveKit (createScreenVideoTrack). Переключение layout VideoGrid на режим «presenter» при активном screenshare. Остановка по кнопке или закрытию вкладки.",
     16, 40_000, ""),

    ("Phase 2", "2.4", "Glossary UI",
     "Страница управления глоссарием: таблица терминов (src_lang / tgt_lang / original / replacement), пагинация и поиск. Импорт CSV / Excel (drag-and-drop). История версий с откатом. Форма добавления / редактирования термина.",
     28, 70_000, "Зависит от Glossary CRUD API"),

    ("Phase 2", "2.5", "Admin Portal — дашборд",
     "React-дашборд: список активных комнат (участники, языки, длительность). GPU загрузка и latency в реальном времени через WebSocket. Графики (Chart.js или recharts). Алерты при превышении порогов.",
     40, 100_000, "Зависит от Admin API (WebSocket метрики)"),

    ("Phase 2", "2.6", "Admin Portal — управление",
     "Управление пользователями: CRUD, смена роли, блокировка. Управление глоссариями (список, версии). Настройки системы (языки, пороги VAD, white-label). История сессий. Архив записей с фильтрами и скачиванием.",
     36, 90_000, ""),

    ("Phase 2", "2.7", "Recording UI",
     "Кнопка запуска/остановки записи в комнате (модератор). Отображение статуса записи. Страница архива записей: список, метаданные (комната, участники, длительность), скачивание по языкам.",
     16, 40_000, "Зависит от Recording API"),
]

# ─────────────────────────────────────────────────────────────────
# ДАННЫЕ: BACKEND
# ─────────────────────────────────────────────────────────────────
# (phase, id, module, задачи, роль, часы, стоимость, примечание)
BACKEND = [
    # ── Phase 1: MVP Backend ──────────────────────────────────────
    ("Phase 1", "1.1", "Инфраструктура (Docker Compose)",
     "Docker Compose стек: livekit/livekit-server, syncvoice (FastAPI + GPU), agent, frontend, nginx (reverse proxy + TLS), certbot (auto-renewal). GPU runtime для syncvoice-контейнера. Настройка livekit.yaml (ports 7880/7881, UDP 50000-50100).",
     "DevOps", 16, 40_000,
     "docker-compose.yml + livekit.yaml уже есть в репо"),

    ("Phase 1", "1.2", "TLS / nginx",
     "nginx.conf: upstream для syncvoice (8000), livekit (7880), frontend. WebSocket proxypass для /lk/. SSL терминация, Let's Encrypt через certbot, auto-renewal каждые 12h.",
     "DevOps", 8, 20_000,
     "Настроено на demomashine.jakeberrimor.com"),

    ("Phase 1", "1.3", "Translation Service (main.py)",
     "FastAPI-сервис: загрузка SeamlessM4Tv2 large при старте (CUDA, FP16). POST /translate: принимает PCM float32, src_lang/tgt_lang query-param, возвращает PCM float32. Поддержка языков: en/ru/hi/zh -> Seamless codes eng/rus/hin/cmn. num_beams=3, speaker_id param.",
     "ML-инженер", 24, 72_000,
     "main.py в репо, модель facebook/seamless-m4t-v2-large"),

    ("Phase 1", "1.4", "Token API (main.py)",
     "GET /token: генерация LiveKit JWT (PyJWT). Параметры: room, identity, lang. Валидация lang (en/ru/hi/zh). Payload: roomJoin, canPublish, canSubscribe, metadata=lang. TTL 1 час.",
     "Backend", 6, 18_000,
     "Реализован в main.py"),

    ("Phase 1", "1.5", "LiveKit Agent (agent.py)",
     "Агент на чистом livekit-rtc SDK (без Agents SDK): watch_rooms() — опрос каждые 5с, создание RoomAgent на новую комнату. RoomAgent: подписка на track_subscribed / participant_disconnected. Динамическое управление TTS-треками (AudioSource + LocalAudioTrack публикация / unpublish).",
     "ML-инженер", 24, 72_000,
     "agent.py в репо"),

    ("Phase 1", "1.6", "VAD + Chunking (agent.py)",
     "RMS-VAD: SPEECH_RMS=0.05 начало, SILENCE_RMS=0.02 продолжение. Подтверждение речи SPEECH_CONFIRM=3 фрейма подряд. Pre-buffer до подтверждения. SILENCE_FRAMES=50 (~1с тишины -> flush). MAX_FRAMES=500 (~10с -> flush без паузы). MIN_SPEECH_SAMPLES=12800 (800мс минимум).",
     "ML-инженер", 16, 48_000,
     "В agent.py _handle_audio()"),

    ("Phase 1", "1.7", "Fanout + Echo Suppression (agent.py)",
     "Fanout: _flush() собирает set tgt_langs из metadata участников (исключая себя и sv-agent). Параллельная трансляция через asyncio.gather. Echo cooldown: после TTS на язык — 2.5с блокировка входящего аудио того же языка. Фильтрация sv_tts_* треков в on_track.",
     "Backend", 12, 36_000,
     "В agent.py _flush() + _translate_and_play()"),

    ("Phase 1", "1.8", "Hallucination Guard (agent.py)",
     "Проверка длительности вывода: если out_dur > max(in_dur * 3, 30.0) — дискард. Логирование предупреждений. Обработка ошибок /translate (status != 200).",
     "ML-инженер", 6, 18_000,
     "В agent.py _translate_and_play()"),

    ("Phase 1", "1.9", "Legacy WebSocket endpoint (main.py)",
     "WebSocket /ws/{room_id}/{user_lang}/{speaker_id}: поддержка старого MVP (браузер без LiveKit). Broadcast PCM от одного пользователя всем остальным с переводом через SeamlessM4T.",
     "Backend", 8, 24_000,
     "В main.py, оставлен для совместимости"),

    # ── Phase 2: Продукт ──────────────────────────────────────────
    ("Phase 2", "2.1", "PostgreSQL Schema",
     "Миграции (Alembic): таблицы users (id, email, name, role, password_hash, created_at, is_active), rooms (session history), glossary_terms (src_lang, tgt_lang, original, replacement, version), glossary_versions, recordings (room_id, path, langs, duration, created_at), audit_log.",
     "Backend", 12, 36_000,
     ""),

    ("Phase 2", "2.2", "Auth API (JWT + роли)",
     "POST /auth/register, /auth/login -> JWT access+refresh токены. Роли: Admin, Moderator, Participant, Guest. CRUD пользователей (Admin only). Гостевой токен по invite-ссылке. Middleware для проверки ролей в FastAPI (Depends).",
     "Backend", 40, 120_000,
     ""),

    ("Phase 2", "2.3", "GlossaryAI (post-processing)",
     "Модуль постобработки текста перед TTS: trie или hash-lookup терминов, case-insensitive замена, правила 'не переводить' (pass-through в кавычках), фонетические транскрипции. Интеграция в /translate после SeamlessM4T generate: text -> replace -> TTS. Поддержка многоязычных глоссариев (ru-en, ru-hi, ru-zh).",
     "ML-инженер", 48, 144_000,
     "Критично для фармы — много спецтерминов"),

    ("Phase 2", "2.4", "Glossary CRUD API",
     "REST: GET/POST/PUT/DELETE /glossary/terms. Фильтрация по src_lang/tgt_lang. POST /glossary/import (CSV, Excel). Версионирование (каждый импорт — новая версия). GET /glossary/versions, POST /glossary/rollback/{version_id}. PostgreSQL backend.",
     "Backend", 28, 84_000,
     ""),

    ("Phase 2", "2.5", "Recording Service",
     "LiveKit Egress API: запуск composite layout записи (активный спикер + grid). Отдельный захват TTS-дорожек по языкам. FFmpeg mux в многодорожечный MP4 (каждый язык — отдельный audio track). Запуск/остановка по REST API (POST /recording/start, /recording/stop).",
     "Backend", 48, 144_000,
     "Клиент: интересует, зависит от цены"),

    ("Phase 2", "2.6", "Recording Storage",
     "Файловое хранилище: сохранение MP4 в /data/recordings/{room_id}/{date}/. Политика retention (configurable, default 30 дней). GET /recording/list (фильтры: room, date, lang). GET /recording/download/{id}. Метаданные в PostgreSQL.",
     "Backend", 16, 48_000,
     ""),

    ("Phase 2", "2.7", "Admin API",
     "GET /admin/rooms — активные комнаты (участники, языки, uptime). WebSocket /admin/metrics — GPU utilization, VRAM, latency по комнатам (push каждые 2с). GET /admin/audit — лог событий (join/leave/translate/error). PUT /admin/settings — пороги VAD, языки, white-label. Доступ только Admin.",
     "Backend", 32, 96_000,
     ""),

    # ── Phase 3: Production-ready ─────────────────────────────────
    ("Phase 3", "3.1", "Мониторинг",
     "Prometheus endpoint /metrics в FastAPI (gpu_utilization, vram_used, translate_latency_ms, active_rooms, active_participants). Grafana дашборд или интеграция в Admin Portal. Алерты: latency > 3с, GPU > 90%, ошибки > 5/мин.",
     "DevOps", 20, 50_000,
     ""),

    ("Phase 3", "3.2", "Production Docker Compose",
     "Финальный compose: health checks для всех сервисов, restart policies, logging (journald или JSON-file с ротацией), env-конфигурация через .env файл. Секреты через docker secrets или env. Отдельные compose для dev/staging/prod.",
     "DevOps", 24, 60_000,
     ""),

    ("Phase 3", "3.3", "Деплой на клиентское железо",
     "Инструкции и скрипты деплоя для: сервера с NVIDIA A100/A10, RTX 4090/5090, Mac Studio (MPS backend, launchd), DGX Spark (DGX OS). GPU runtime тест. install.sh с проверками зависимостей.",
     "DevOps", 40, 100_000,
     "4 конфигурации железа"),

    ("Phase 3", "3.4", "Документация — API",
     "OpenAPI spec (FastAPI /docs автогенерация + ручная правка). Примеры интеграции для клиента: подключение к LiveKit, получение токена, смена языка. Postman-коллекция.",
     "Backend", 12, 36_000,
     ""),

    ("Phase 3", "3.5", "Документация — пользователь",
     "Руководство модератора: создание комнаты, управление участниками, запись. Руководство участника: вход, выбор языка, управление микро/камерой. Руководство администратора: глоссарии, метрики, пользователи.",
     "Backend", 16, 48_000,
     ""),

    ("Phase 3", "3.6", "Приёмочное тестирование",
     "Чеклист по всем FR/NFR из ТЗ. E2E тест: 10-20 участников, 3 языка, глоссарий, запись, admin portal. Latency замеры на production-железе. Баг-фиксы. Финальная стабилизация.",
     "QA", 56, 112_000,
     ""),
]

# ─────────────────────────────────────────────────────────────────
# СТИЛИ
# ─────────────────────────────────────────────────────────────────

def thin():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)

PHASE_COLOR = {
    "Phase 1": "D6EAF8",
    "Phase 2": "D5F5E3",
    "Phase 3": "FDEBD0",
}
ROLE_COLOR = {
    "ML-инженер": "F3E8FF",
    "Backend":    "EAF4FF",
    "DevOps":     "FFFBE6",
    "QA":         "FDEDEC",
    "Frontend":   "E8F8F5",
}

def fill(hex_c):
    return PatternFill("solid", fgColor=hex_c)

def hdr(ws, row, ncols, bg, fg="FFFFFF", size=11):
    for c in range(1, ncols+1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill(bg)
        cell.font = Font(bold=True, color=fg, size=size)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin()

def phase_row(ws, row, phase, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1)
    c.value = f"  ▸ {phase}"
    c.fill = fill(PHASE_COLOR.get(phase, "EEEEEE"))
    c.font = Font(bold=True, size=10, color="2C3E50")
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = thin()
    ws.row_dimensions[row].height = 20

def data_cell(cell, center=False, bold=False, bg=None, fg=None):
    cell.alignment = Alignment(horizontal="center" if center else "left",
                               vertical="center", wrap_text=True)
    cell.border = thin()
    if bg: cell.fill = fill(bg)
    if fg or bold: cell.font = Font(bold=bold, color=fg or "000000")

# ─────────────────────────────────────────────────────────────────
# ЛИСТ 1: FRONTEND
# ─────────────────────────────────────────────────────────────────

ws_f = wb.active
ws_f.title = "Frontend"
ws_f.sheet_view.showGridLines = False

F_COLS  = ["Фаза", "#", "Модуль", "Описание задачи", "Часы", "Стоимость", "Примечание"]
F_WIDTHS = [12, 5, 20, 60, 7, 14, 30]

# Заголовок листа
ws_f.row_dimensions[1].height = 38
ws_f.merge_cells("A1:G1")
c = ws_f["A1"]
c.value = "SyncVoice™  —  Frontend"
c.fill = fill("1A3A5C"); c.font = Font(bold=True, size=14, color="FFFFFF")
c.alignment = Alignment(horizontal="center", vertical="center")
c.border = thin()

# Шапка колонок
ws_f.row_dimensions[2].height = 30
hdr(ws_f, 2, len(F_COLS), "2980B9")
for i, h in enumerate(F_COLS, 1):
    ws_f.cell(row=2, column=i, value=h)

row = 3
cur_phase = None
for task in FRONTEND:
    phase, tid, module, desc, hours, cost, notes = task
    if phase != cur_phase:
        phase_row(ws_f, row, phase, len(F_COLS))
        row += 1
        cur_phase = phase

    ws_f.row_dimensions[row].height = 60
    vals = [phase, tid, module, desc, hours, f"{cost:,} ₽".replace(",", " "), notes]
    for ci, v in enumerate(vals, 1):
        c = ws_f.cell(row=row, column=ci, value=v)
        data_cell(c,
                  center=(ci in [1, 2, 5, 6]),
                  bold=(ci == 3),
                  bg=(PHASE_COLOR.get(phase, "FFFFFF") if ci == 1 else None))
    row += 1

# Итого
ws_f.row_dimensions[row].height = 28
total_h = sum(t[4] for t in FRONTEND)
total_c = sum(t[5] for t in FRONTEND)
ws_f.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
c = ws_f.cell(row=row, column=1, value="ИТОГО  Frontend")
c.fill = fill("1A3A5C"); c.font = Font(bold=True, color="FFFFFF")
c.alignment = Alignment(horizontal="center", vertical="center"); c.border = thin()
for ci, v in [(5, total_h), (6, f"{total_c:,} ₽".replace(",", " ")), (7, "")]:
    c2 = ws_f.cell(row=row, column=ci, value=v)
    c2.fill = fill("1A3A5C"); c2.font = Font(bold=True, color="FFFFFF")
    c2.alignment = Alignment(horizontal="center", vertical="center"); c2.border = thin()

for i, w in enumerate(F_WIDTHS, 1):
    ws_f.column_dimensions[get_column_letter(i)].width = w

# ─────────────────────────────────────────────────────────────────
# ЛИСТ 2: BACKEND
# ─────────────────────────────────────────────────────────────────

ws_b = wb.create_sheet("Backend")
ws_b.sheet_view.showGridLines = False

B_COLS  = ["Фаза", "#", "Модуль", "Описание задачи", "Роль", "Часы", "Стоимость", "Примечание"]
B_WIDTHS = [12, 5, 22, 56, 13, 7, 14, 28]

ws_b.row_dimensions[1].height = 38
ws_b.merge_cells("A1:H1")
c = ws_b["A1"]
c.value = "SyncVoice™  —  Backend"
c.fill = fill("1B3A2A"); c.font = Font(bold=True, size=14, color="FFFFFF")
c.alignment = Alignment(horizontal="center", vertical="center")
c.border = thin()

ws_b.row_dimensions[2].height = 30
hdr(ws_b, 2, len(B_COLS), "1E8449")
for i, h in enumerate(B_COLS, 1):
    ws_b.cell(row=2, column=i, value=h)

row = 3
cur_phase = None
for task in BACKEND:
    phase, tid, module, desc, role, hours, cost, notes = task
    if phase != cur_phase:
        phase_row(ws_b, row, phase, len(B_COLS))
        row += 1
        cur_phase = phase

    ws_b.row_dimensions[row].height = 68
    vals = [phase, tid, module, desc, role, hours, f"{cost:,} ₽".replace(",", " "), notes]
    for ci, v in enumerate(vals, 1):
        c = ws_b.cell(row=row, column=ci, value=v)
        data_cell(c,
                  center=(ci in [1, 2, 5, 6, 7]),
                  bold=(ci == 3),
                  bg=(PHASE_COLOR.get(phase, "FFFFFF") if ci == 1
                      else ROLE_COLOR.get(role, "FFFFFF") if ci == 5
                      else None))
    row += 1

# Итого
ws_b.row_dimensions[row].height = 28
total_h = sum(t[5] for t in BACKEND)
total_c = sum(t[6] for t in BACKEND)
ws_b.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
c = ws_b.cell(row=row, column=1, value="ИТОГО  Backend / DevOps / ML / QA")
c.fill = fill("1B3A2A"); c.font = Font(bold=True, color="FFFFFF")
c.alignment = Alignment(horizontal="center", vertical="center"); c.border = thin()
for ci, v in [(5, ""), (6, total_h), (7, f"{total_c:,} ₽".replace(",", " ")), (8, "")]:
    c2 = ws_b.cell(row=row, column=ci, value=v)
    c2.fill = fill("1B3A2A"); c2.font = Font(bold=True, color="FFFFFF")
    c2.alignment = Alignment(horizontal="center", vertical="center"); c2.border = thin()

for i, w in enumerate(B_WIDTHS, 1):
    ws_b.column_dimensions[get_column_letter(i)].width = w

# ─────────────────────────────────────────────────────────────────
out = "/home/clawdbot/.openclaw/workspace/drafts/SyncVoice_Decomp_Front_Back_v2.xlsx"
wb.save(out)
print(f"Saved: {out}")
print(f"Frontend: {len(FRONTEND)} tasks, {sum(t[4] for t in FRONTEND)}h, {sum(t[5] for t in FRONTEND):,}₽")
print(f"Backend:  {len(BACKEND)} tasks, {sum(t[5] for t in BACKEND)}h, {sum(t[6] for t in BACKEND):,}₽")
