#!/usr/bin/env python3
"""
Генератор ЗАЯВОК для Gravity Group (с ЦВЕТНЫМИ ГРАНИЦАМИ таблиц)
Использование: python3 gravity_zavka_template_v3.py
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from datetime import datetime

def set_table_borders_colored(table, color_hex='5F497A', inside_only=False):
    """
    Установка цветных границ таблицы (как в оригинале Gravity Group)
    
    Args:
        table: Таблица docx
        color_hex: Цвет границ в hex (по умолчанию 5F497A - фиолетовый)
        inside_only: Если True, красить только внутренние границы
    """
    tbl = table._element
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    
    # Создаём элемент границ
    tblBorders = OxmlElement('w:tblBorders')
    
    # Определяем какие границы красить
    if inside_only:
        borders_to_set = ['insideH', 'insideV']
    else:
        borders_to_set = ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']
    
    for border_name in borders_to_set:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')  # Толщина
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color_hex)
        tblBorders.append(border)
    
    # Удаляем старые границы если есть
    old_borders = tblPr.find(qn('w:tblBorders'))
    if old_borders is not None:
        tblPr.remove(old_borders)
    
    tblPr.append(tblBorders)

def number_to_russian_words(num):
    """Конвертация числа в слова на русском языке"""
    if num == 0:
        return "ноль"
    
    ones = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    teens = ["десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать", 
             "пятнадцать", "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать"]
    tens = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", 
            "шестьдесят", "семьдесят", "восемьдесят", "девяносто"]
    hundreds = ["", "сто", "двести", "триста", "четыреста", "пятьсот", 
                "шестьсот", "семьсот", "восемьсот", "девятьсот"]
    
    thousands = ["", "одна", "две", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    
    def convert_hundreds(n):
        if n == 0:
            return ""
        elif n < 10:
            return ones[n]
        elif n < 20:
            return teens[n - 10]
        elif n < 100:
            return tens[n // 10] + (" " + ones[n % 10] if n % 10 != 0 else "")
        else:
            h = hundreds[n // 100]
            remainder = n % 100
            if remainder == 0:
                return h
            elif remainder < 10:
                return h + " " + ones[remainder]
            elif remainder < 20:
                return h + " " + teens[remainder - 10]
            else:
                return h + " " + tens[remainder // 10] + (" " + ones[remainder % 10] if remainder % 10 != 0 else "")
    
    def get_thousand_form(n):
        if n % 10 == 1 and n % 100 != 11:
            return "тысяча"
        elif n % 10 in [2, 3, 4] and n % 100 not in [12, 13, 14]:
            return "тысячи"
        else:
            return "тысяч"
    
    num = int(num)
    if num == 0:
        return "ноль"
    
    result = []
    
    # Миллионы
    millions = num // 1000000
    if millions > 0:
        result.append(convert_hundreds(millions))
        if millions % 10 == 1 and millions % 100 != 11:
            result.append("миллион")
        elif millions % 10 in [2, 3, 4] and millions % 100 not in [12, 13, 14]:
            result.append("миллиона")
        else:
            result.append("миллионов")
    
    # Тысячи
    thousands_num = (num % 1000000) // 1000
    if thousands_num > 0:
        if thousands_num < 10:
            result.append(thousands[thousands_num])
        elif thousands_num < 20:
            result.append(teens[thousands_num - 10])
        elif thousands_num < 100:
            t = tens[thousands_num // 10]
            o = thousands[thousands_num % 10] if thousands_num % 10 != 0 else ""
            result.append(t + (" " + o if o else ""))
        else:
            h = hundreds[thousands_num // 100]
            remainder = thousands_num % 100
            if remainder == 0:
                result.append(h)
            elif remainder < 10:
                result.append(h + " " + thousands[remainder])
            elif remainder < 20:
                result.append(h + " " + teens[remainder - 10])
            else:
                t = tens[remainder // 10]
                o = thousands[remainder % 10] if remainder % 10 != 0 else ""
                result.append(h + " " + t + (" " + o if o else ""))
        
        result.append(get_thousand_form(thousands_num))
    
    # Единицы
    ones_num = num % 1000
    if ones_num > 0:
        result.append(convert_hundreds(ones_num))
    
    return " ".join(result)

def format_money(amount):
    """
    Форматирование денег в стиле оригинала: 147 450 (пробел как разделитель)
    """
    rubles = int(amount)
    kopeks = int(round((amount - rubles) * 100))
    
    # Форматируем рубли с пробелами как разделителями тысяч
    rubles_str = f"{rubles:,}".replace(',', ' ')
    
    return rubles_str, kopeks

def money_to_words(amount):
    """Конвертация суммы в слова"""
    rubles = int(amount)
    kopeks = int(round((amount - rubles) * 100))
    
    rubles_words = number_to_russian_words(rubles).capitalize()
    
    return f"{rubles_words} рублей, {kopeks:02d} копеек"

def create_zavka_document(data):
    """
    Создание ТОЧНОЙ КОПИИ заявки Gravity Group (с цветными границами)
    """
    
    doc = Document()
    
    # === ЗАГОЛОВОК (БЕЗ ЦЕНТРИРОВАНИЯ!) ===
    title = doc.add_paragraph()
    run = title.add_run(f"ЗАЯВКА №{data['zavka_number']} к Договору {data['contract_number']} от «{data['contract_date']}» г.")
    run.font.bold = True
    run.font.size = Pt(12)
    
    # Дата заявки
    date_para = doc.add_paragraph()
    run = date_para.add_run(f"Дата Заявки: «{data['zavka_date']}» г.")
    run.font.bold = True
    run.font.size = Pt(12)
    
    # === ТАБЛИЦА СТОРОН ===
    table_parties = doc.add_table(rows=2, cols=2)
    table_parties.style = 'Normal Table'
    set_table_borders_colored(table_parties, color_hex='5F497A', inside_only=True)
    
    # Заголовки
    cell = table_parties.rows[0].cells[0]
    cell.text = "Заказчик"
    
    cell = table_parties.rows[0].cells[1]
    customer_text = (
        f"{data['customer_name']}\n"
        f"ИНН {data['customer_inn']}\n"
        f"в лице директора {data['customer_director']}\n"
        f"действующего на основании Устава"
    )
    cell.text = customer_text
    
    # Исполнитель
    cell = table_parties.rows[1].cells[0]
    cell.text = "Исполнитель"
    
    cell = table_parties.rows[1].cells[1]
    executor_text = (
        "Общество с ограниченной ответственностью «Гравити Групп»\n"
        "ИНН 5908999996\n"
        "в лице Директора Яборова Андрея Владимировича, действующего на основании Устава"
    )
    cell.text = executor_text
    
    # Шрифт 12pt во всех ячейках
    for row in table_parties.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(12)
    
    # === ОСНОВНОЙ ТЕКСТ (БЕЗ ОТСТУПОВ!) ===
    para = doc.add_paragraph("согласовали следующий условия выполнения работ по Заявке к Договору:")
    for run in para.runs:
        run.font.size = Pt(12)
    
    # Стоимость
    rubles_str, kopeks = format_money(data['cost_total'])
    cost_text = f"Предварительная стоимость работ по Заявке составляет {rubles_str} ({money_to_words(data['cost_total'])})"
    if data.get('cost_vat_percent'):
        vat_amount = data['cost_total'] * data['cost_vat_percent'] / (100 + data['cost_vat_percent'])
        vat_str, vat_kop = format_money(vat_amount)
        cost_text += f", в т.ч. НДС {data['cost_vat_percent']}% - {vat_str},{vat_kop:02d} рублей"
    cost_text += "."
    
    para = doc.add_paragraph(cost_text)
    for run in para.runs:
        run.font.size = Pt(12)
    
    # Сроки
    para = doc.add_paragraph("Сроки выполнения Заявки:")
    for run in para.runs:
        run.font.size = Pt(12)
    
    para = doc.add_paragraph(f"Дата начала выполнения работ: «{data['work_start']}» г.")
    for run in para.runs:
        run.font.size = Pt(12)
    
    para = doc.add_paragraph(f"Дата окончания выполнения работ: «{data['work_end']}» г.")
    for run in para.runs:
        run.font.size = Pt(12)
    
    # Порядок оплаты
    para = doc.add_paragraph("Порядок оплаты работ по Заявке:")
    for run in para.runs:
        run.font.size = Pt(12)
    
    pay30_str, pay30_kop = format_money(data['payment_30_percent'])
    para = doc.add_paragraph(
        f"30% планируемой стоимости работ в размере {pay30_str} ({money_to_words(data['payment_30_percent'])}) "
        f"оплачиваются в течение 2 (Двух) рабочих с момента подписания Заявки."
    )
    for run in para.runs:
        run.font.size = Pt(12)
    
    pay37_str, pay37_kop = format_money(data['payment_37_percent'])
    para = doc.add_paragraph(
        f"37% планируемой стоимости работ в размере {pay37_str} ({money_to_words(data['payment_37_percent'])}) "
        f"оплачиваются в течение 30 (Тридцати) рабочих с момента подписания Заявки, но не позже даты подписания Акта по Заявке."
    )
    for run in para.runs:
        run.font.size = Pt(12)
    
    para = doc.add_paragraph(
        "Оплата оставшейся части общей стоимости работ, рассчитанной на основании фактических "
        "трудозатрат Исполнителя, и указанных в Акте по производится в течение 5 (Пяти) рабочих "
        "дней со дня подписания Акта по Заявке. По согласованию с Исполнителем оплата оставшейся "
        "части общей стоимости выполненных работ может быть отсрочена до 20 (Двадцати) рабочих "
        "дней после подписания Акта сдачи-приемки работ."
    )
    for run in para.runs:
        run.font.size = Pt(12)
    
    # Гарантия
    warranty_word = number_to_russian_words(data['warranty_months']).capitalize()
    para = doc.add_paragraph(
        f"Срок гарантии на выполненные работы Исполнителем составляет {data['warranty_months']} "
        f"({warranty_word}) месяца."
    )
    for run in para.runs:
        run.font.size = Pt(12)
    
    # Доп условие
    para = doc.add_paragraph(
        "Если для выполнения Заявки потребуются проработка предметной области, разработка 3D моделей, "
        "то для выполнения работ со стороны Исполнителя привлекается аналитик, разработчик, "
        "трудозатраты будут отражены в Акте по Заявке."
    )
    for run in para.runs:
        run.font.size = Pt(12)
    
    # Приоритет
    para = doc.add_paragraph(
        "Положения Заявки имеют приоритет над условиями Договора. В остальном, что не предусмотрено "
        "Заявкой, Стороны руководствуются условиями Договора."
    )
    for run in para.runs:
        run.font.size = Pt(12)
    
    # === ТАБЛИЦА РАБОТ ===
    # Подсчёт строк
    total_rows = 1  # Заголовок
    for work_cat in data['works']:
        total_rows += 1  # Категория
        total_rows += len(work_cat['items'])  # Специалисты
    
    table_works = doc.add_table(rows=total_rows, cols=3)
    table_works.style = 'Normal Table'
    set_table_borders_colored(table_works, color_hex='5F497A', inside_only=True)
    
    # Заголовок
    headers = ['№', 'Наименование и описание работ', 'Оценка трудоёмкости в часах']
    for i, header in enumerate(headers):
        cell = table_works.rows[0].cells[i]
        cell.text = header
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(12)
    
    # Заполнение
    row_idx = 1
    for work_cat in data['works']:
        cell = table_works.rows[row_idx].cells[1]
        cell.text = work_cat['category']
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(12)
        row_idx += 1
        
        for item in work_cat['items']:
            table_works.rows[row_idx].cells[1].text = item['specialist']
            # Формат часов: 23,00 (запятая как разделитель!)
            table_works.rows[row_idx].cells[2].text = f"{item['hours']:.2f}".replace('.', ',')
            
            for cell in table_works.rows[row_idx].cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(12)
            row_idx += 1
    
    # === ТАБЛИЦА ПОДПИСЕЙ ===
    table_signs = doc.add_table(rows=4, cols=2)
    table_signs.style = 'Normal Table'
    set_table_borders_colored(table_signs, color_hex='5F497A', inside_only=False)
    
    # Заголовки
    cell = table_signs.rows[0].cells[0]
    cell.text = "ЗАКАЗЧИК"
    cell.paragraphs[0].runs[0].font.bold = True
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    cell = table_signs.rows[0].cells[1]
    cell.text = "ИСПОЛНИТЕЛЬ"
    cell.paragraphs[0].runs[0].font.bold = True
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Организации
    table_signs.rows[1].cells[0].text = ""
    table_signs.rows[1].cells[1].text = "Общество с ограниченной ответственностью «Гравити Групп»"
    
    # Должности
    cell = table_signs.rows[2].cells[0]
    cell.text = (
        f"Генеральный директор {data['customer_name']}\n"
        f"{data['customer_director']}\n\n"
        f"(на основании Устава)"
    )
    
    cell = table_signs.rows[2].cells[1]
    cell.text = (
        "Директор ООО «ГРАВИТИ ГРУПП»\n"
        "Яборов Андрей Владимирович\n"
        "(на основании Устава)"
    )
    
    # Подписи
    cell = table_signs.rows[3].cells[0]
    cell.text = f"_____________________ / {data['customer_director_short']}"
    
    cell = table_signs.rows[3].cells[1]
    cell.text = "_________________________ / А.В. Яборов"
    
    # Шрифт 12pt везде
    for row in table_signs.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(12)
    
    return doc

# === ПРИМЕР ===
if __name__ == "__main__":
    example_data = {
        'zavka_number': '4',
        'contract_number': '14/25',
        'contract_date': '16» мая 2025',  # Обратите внимание на формат!
        'zavka_date': '05» февраля 2026',
        'cost_total': 166057.50,
        'cost_vat_percent': 5,
        'work_start': '5» февраля 2026',
        'work_end': '27» февраля 2026',
        'payment_30_percent': 49817.25,
        'payment_37_percent': 61441.28,
        'warranty_months': 3,
        'customer_name': 'ООО "Инплайн"',
        'customer_inn': '5902054490',
        'customer_director': 'Шевшелев Сергей Викторович',
        'customer_director_short': 'С.В. Шевшелев',
        'works': [
            {'category': 'Разработка мобильной версии приложения (Приложение №1 к Заявке)', 'items': [
                {'specialist': 'Дизайнер', 'hours': 8.0},
                {'specialist': 'Веб разработчик', 'hours': 23.0},
                {'specialist': 'Специалист по тестированию', 'hours': 10.0}
            ]}
        ]
    }
    
    doc = create_zavka_document(example_data)
    
    output_path = "/home/clawdbot/.openclaw/workspace/docs/gravity_zavka_v3_COLORED.docx"
    doc.save(output_path)
    
    print(f"✅ Шаблон с цветными границами создан: {output_path}")
    print("\n📋 ИСПРАВЛЕНИЯ:")
    print("  ✅ Цветные границы таблиц (5F497A - фиолетовый)")
    print("  ✅ inside_only для таблиц сторон и работ")
    print("  ✅ Все границы для таблицы подписей")
