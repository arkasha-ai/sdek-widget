#!/usr/bin/env python3
"""
Детальный анализ ВСЕХ docx документов Gravity Group
Извлекает мельчайшие детали форматирования
"""

from docx import Document
from docx.shared import Pt, Cm, Inches
from lxml import etree
import os

def analyze_font(run):
    """Анализ шрифта run"""
    details = {}
    if run.font.name:
        details['name'] = run.font.name
    if run.font.size:
        details['size'] = f"{run.font.size.pt}pt"
    if run.bold:
        details['bold'] = True
    if run.italic:
        details['italic'] = True
    if run.underline:
        details['underline'] = True
    if run.font.color and run.font.color.rgb:
        details['color'] = str(run.font.color.rgb)
    return details

def analyze_paragraph(para):
    """Анализ параграфа"""
    details = {}
    
    # Выравнивание
    if para.alignment:
        details['alignment'] = str(para.alignment)
    
    # Отступы
    pf = para.paragraph_format
    if pf.left_indent:
        details['left_indent'] = f"{pf.left_indent.pt:.2f}pt"
    if pf.right_indent:
        details['right_indent'] = f"{pf.right_indent.pt:.2f}pt"
    if pf.first_line_indent:
        details['first_line_indent'] = f"{pf.first_line_indent.pt:.2f}pt"
    
    # Интервалы
    if pf.space_before:
        details['space_before'] = f"{pf.space_before.pt:.2f}pt"
    if pf.space_after:
        details['space_after'] = f"{pf.space_after.pt:.2f}pt"
    if pf.line_spacing:
        details['line_spacing'] = str(pf.line_spacing)
    
    # Стиль
    if para.style and para.style.name:
        details['style'] = para.style.name
    
    return details

def analyze_table_borders(table):
    """Анализ границ таблицы"""
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    
    tbl = table._element
    tblPr = tbl.find('.//w:tblPr', ns)
    
    borders = {}
    if tblPr is not None:
        tblBorders = tblPr.find('.//w:tblBorders', ns)
        if tblBorders is not None:
            for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                border = tblBorders.find(f'.//w:{border_name}', ns)
                if border is not None:
                    color = border.get('{' + ns['w'] + '}color')
                    val = border.get('{' + ns['w'] + '}val')
                    sz = border.get('{' + ns['w'] + '}sz')
                    if color or val:
                        borders[border_name] = {'color': color, 'val': val, 'sz': sz}
    
    return borders

def analyze_cell_borders(cell):
    """Анализ границ ячейки"""
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    
    tc = cell._element
    tcPr = tc.find('.//w:tcPr', ns)
    
    borders = {}
    if tcPr is not None:
        tcBorders = tcPr.find('.//w:tcBorders', ns)
        if tcBorders is not None:
            for border_name in ['top', 'left', 'bottom', 'right']:
                border = tcBorders.find(f'.//w:{border_name}', ns)
                if border is not None:
                    color = border.get('{' + ns['w'] + '}color')
                    val = border.get('{' + ns['w'] + '}val')
                    sz = border.get('{' + ns['w'] + '}sz')
                    if color or val:
                        borders[border_name] = {'color': color, 'val': val, 'sz': sz}
    
    return borders

def analyze_document(file_path):
    """Полный анализ документа"""
    print(f"\n{'='*80}")
    print(f"ФАЙЛ: {os.path.basename(file_path)}")
    print(f"{'='*80}\n")
    
    doc = Document(file_path)
    
    # === СЕКЦИИ И MARGINS ===
    print("📐 СЕКЦИИ И MARGINS:")
    for i, section in enumerate(doc.sections):
        print(f"\n  Секция {i+1}:")
        print(f"    Ориентация: {section.orientation}")
        print(f"    Размер страницы: {section.page_width.cm:.2f}см x {section.page_height.cm:.2f}см")
        print(f"    Margins:")
        print(f"      Top: {section.top_margin.cm:.2f}см")
        print(f"      Bottom: {section.bottom_margin.cm:.2f}см")
        print(f"      Left: {section.left_margin.cm:.2f}см")
        print(f"      Right: {section.right_margin.cm:.2f}см")
    
    # === СТИЛИ ДОКУМЕНТА ===
    print("\n\n🎨 СТИЛИ ДОКУМЕНТА:")
    styles_used = set()
    for para in doc.paragraphs:
        if para.style and para.style.name:
            styles_used.add(para.style.name)
    for table in doc.tables:
        if table.style and table.style.name:
            styles_used.add(table.style.name)
    
    for style_name in sorted(styles_used):
        print(f"  • {style_name}")
    
    # === ПАРАГРАФЫ (первые 10) ===
    print("\n\n📝 ПАРАГРАФЫ (первые 10):")
    for i, para in enumerate(doc.paragraphs[:10]):
        if para.text.strip():
            print(f"\n  [{i+1}] {para.text[:60]}{'...' if len(para.text) > 60 else ''}")
            
            # Детали параграфа
            para_details = analyze_paragraph(para)
            if para_details:
                for key, val in para_details.items():
                    print(f"      {key}: {val}")
            
            # Детали runs
            if para.runs:
                print(f"      Runs: {len(para.runs)}")
                for j, run in enumerate(para.runs[:3]):  # Первые 3 run
                    if run.text.strip():
                        run_details = analyze_font(run)
                        if run_details:
                            print(f"        Run {j+1}: {run_details}")
    
    # === ТАБЛИЦЫ ===
    print(f"\n\n📊 ТАБЛИЦЫ ({len(doc.tables)} шт):")
    for i, table in enumerate(doc.tables):
        print(f"\n  Таблица {i+1}: {len(table.rows)} строк x {len(table.columns)} колонок")
        
        # Стиль таблицы
        if table.style:
            print(f"    Стиль: {table.style.name}")
        
        # Границы таблицы
        borders = analyze_table_borders(table)
        if borders:
            print(f"    Границы таблицы:")
            for border_name, border_info in borders.items():
                print(f"      {border_name}: color={border_info['color']}, val={border_info['val']}, sz={border_info['sz']}")
        
        # Первая ячейка (для примера)
        if table.rows:
            cell = table.rows[0].cells[0]
            cell_borders = analyze_cell_borders(cell)
            if cell_borders:
                print(f"    Границы ячейки [0,0]:")
                for border_name, border_info in cell_borders.items():
                    print(f"      {border_name}: color={border_info['color']}, val={border_info['val']}, sz={border_info['sz']}")
            
            # Текст первой ячейки
            if cell.text.strip():
                print(f"    Ячейка [0,0]: {cell.text[:40]}{'...' if len(cell.text) > 40 else ''}")
                if cell.paragraphs:
                    para = cell.paragraphs[0]
                    para_details = analyze_paragraph(para)
                    if para_details:
                        print(f"      Параграф: {para_details}")
                    if para.runs:
                        run = para.runs[0]
                        run_details = analyze_font(run)
                        if run_details:
                            print(f"      Run: {run_details}")
        
        # Ширина колонок
        print(f"    Ширина колонок:")
        for j, col in enumerate(table.columns):
            if col.width:
                print(f"      Колонка {j+1}: {col.width.cm:.2f}см")

# === АНАЛИЗ ВСЕХ ФАЙЛОВ ===
if __name__ == "__main__":
    files = [
        '/home/clawdbot/.openclaw/media/inbound/file_72---efdbaa49-ab1e-4bf0-87fc-3a50c33684f1.docx',
        '/home/clawdbot/.openclaw/media/inbound/file_73---9375919f-b4fd-4a66-95aa-aee8b6d2a62a.docx',
        '/home/clawdbot/.openclaw/media/inbound/file_74---0e742e31-30cf-47ec-9217-659ff5a9060d.docx',
        '/home/clawdbot/.openclaw/media/inbound/file_75---cc5d8155-951a-4ec8-9264-14e676183e15.docx',
        '/home/clawdbot/.openclaw/media/inbound/file_78---a90fd421-7a8b-48cc-b81c-e0be7c6fdd75.docx',
        '/home/clawdbot/.openclaw/media/inbound/file_80---cd7c35cb-d362-4faa-a878-d39219c9640f.docx'
    ]
    
    for file_path in files:
        if os.path.exists(file_path):
            try:
                analyze_document(file_path)
            except Exception as e:
                print(f"❌ Ошибка при анализе {file_path}: {e}")
        else:
            print(f"❌ Файл не найден: {file_path}")
    
    print(f"\n\n{'='*80}")
    print("✅ АНАЛИЗ ЗАВЕРШЁН")
    print(f"{'='*80}\n")
