import openpyxl
from openpyxl import Workbook

# Load the source workbook with data_only=True to get actual values
source_wb = openpyxl.load_workbook('/home/clawdbot/.openclaw/workspace/drafts/SyncVoice_Decomposition_v1_fixed_final.xlsx', data_only=True)
source_ws = source_wb['Декомпозиция']

# Create a new workbook
new_wb = Workbook()
new_ws = new_wb.active
new_ws.title = 'Декомпозиция'

# Copy data from rows 4-34, columns A-J
for row in range(4, 35):  # 35 because range is exclusive
    for col in range(1, 11):  # 11 because range is exclusive (A-J)
        new_ws.cell(row=row, column=col).value = source_ws.cell(row=row, column=col).value

# Add formulas to column K (Итого часов) for rows 4-34
for row in range(4, 35):
    new_ws.cell(row=row, column=11).value = f'=E{row}+F{row}+G{row}+H{row}+I{row}+J{row}'  # Column K

# Add formulas to column L (Стоимость) for rows 4-34
for row in range(4, 35):
    new_ws.cell(row=row, column=12).value = f'=E{row}*3000+F{row}*3000+G{row}*2500+H{row}*2500+I{row}*2000+J{row}*3500'  # Column L

# Add total formulas in row 36 (ИТОГО без буфера)
new_ws.cell(row=36, column=5).value = '=SUM(E4:E34)'  # Column E
new_ws.cell(row=36, column=6).value = '=SUM(F4:F34)'  # Column F
new_ws.cell(row=36, column=7).value = '=SUM(G4:G34)'  # Column G
new_ws.cell(row=36, column=8).value = '=SUM(H4:H34)'  # Column H
new_ws.cell(row=36, column=9).value = '=SUM(I4:I34)'  # Column I
new_ws.cell(row=36, column=10).value = '=SUM(J4:J34)'  # Column J
new_ws.cell(row=36, column=11).value = '=SUM(K4:K34)'  # Column K
new_ws.cell(row=36, column=12).value = '=SUM(L4:L34)'  # Column L

# Add buffer formulas in row 37 (Буфер 15%)
new_ws.cell(row=37, column=5).value = '=E36*0.15'  # Column E
new_ws.cell(row=37, column=6).value = '=F36*0.15'  # Column F
new_ws.cell(row=37, column=7).value = '=G36*0.15'  # Column G
new_ws.cell(row=37, column=8).value = '=H36*0.15'  # Column H
new_ws.cell(row=37, column=9).value = '=I36*0.15'  # Column I
new_ws.cell(row=37, column=10).value = '=J36*0.15'  # Column J
new_ws.cell(row=37, column=11).value = '=K36*0.15'  # Column K
new_ws.cell(row=37, column=12).value = '=L36*0.15'  # Column L

# Add total with buffer formulas in row 38 (ИТОГО С БУФЕРОМ)
new_ws.cell(row=38, column=5).value = '=E36+E37'  # Column E
new_ws.cell(row=38, column=6).value = '=F36+F37'  # Column F
new_ws.cell(row=38, column=7).value = '=G36+G37'  # Column G
new_ws.cell(row=38, column=8).value = '=H36+H37'  # Column H
new_ws.cell(row=38, column=9).value = '=I36+I37'  # Column I
new_ws.cell(row=38, column=10).value = '=J36+J37'  # Column J
new_ws.cell(row=38, column=11).value = '=K36+K37'  # Column K
new_ws.cell(row=38, column=12).value = '=L36+L37'  # Column L

# Save the new workbook
new_wb.save('/home/clawdbot/.openclaw/workspace/drafts/SyncVoice_Decomposition_v3_final.xlsx')

print("Excel file has been successfully created with proper formulas.")