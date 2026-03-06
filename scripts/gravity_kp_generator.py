"""
КП (Коммерческое предложение) Generator for Gravity Group
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

LOGO_PATH = "/home/clawdbot/.openclaw/workspace/docs/gravity_images/file_74---0e742e31-30cf-47ec-9217-659ff5a9060d_image_1.png"

# Colors
TEAL = RGBColor(0x00, 0x97, 0x91)   # Gravity teal
DARK = RGBColor(0x26, 0x26, 0x26)   # Near black
GRAY = RGBColor(0x60, 0x60, 0x60)   # Medium gray
LIGHT_TEAL_BG = "E8F5F5"             # Very light teal for header bg
TEAL_HEX = "009791"
DARK_HEX = "262626"

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if val is not None:
            b = OxmlElement(f'w:{side}')
            b.set(qn('w:val'), val.get('val', 'single'))
            b.set(qn('w:sz'), str(val.get('sz', 4)))
            b.set(qn('w:color'), val.get('color', '000000'))
            b.set(qn('w:space'), '0')
            tcBorders.append(b)
    tcPr.append(tcBorders)

def set_no_borders(table):
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'nil')
        tblBorders.append(b)
    tblPr.append(tblBorders)

def add_paragraph(doc, text='', bold=False, italic=False, size=11, color=None,
                   align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        run.font.name = 'Calibri'
        if color:
            run.font.color.rgb = color
    return p

def generate_kp(data: dict, output_path: str):
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    # ─────────────── HEADER: Logo + Реквизиты ───────────────
    header_tbl = doc.add_table(rows=1, cols=2)
    header_tbl.autofit = False
    set_no_borders(header_tbl)

    # Ширины колонок (total ~16cm)
    header_tbl.columns[0].width = Cm(8)
    header_tbl.columns[1].width = Cm(8)

    # Лого (левая ячейка)
    logo_cell = header_tbl.cell(0, 0)
    logo_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    logo_para = logo_cell.paragraphs[0]
    logo_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo_run = logo_para.add_run()
    logo_run.add_picture(LOGO_PATH, width=Cm(4.5))

    # Реквизиты (правая ячейка)
    req_cell = header_tbl.cell(0, 1)
    req_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    req_para = req_cell.paragraphs[0]
    req_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    req_para.paragraph_format.space_after = Pt(0)
    for line in [
        ("ООО «Гравити Групп»", True, 9, DARK),
        ("ИНН 5908999996", False, 8, GRAY),
        ("г. Пермь", False, 8, GRAY),
        (data.get("contractor_contact", ""), False, 8, TEAL),
    ]:
        r = req_para.add_run(("\n" if req_para.runs else "") + line[0])
        r.bold = line[1]
        r.font.size = Pt(line[2])
        r.font.name = 'Calibri'
        r.font.color.rgb = line[3]

    doc.add_paragraph()  # небольшой отступ

    # ─────────────── РАЗДЕЛИТЕЛЬ (цветная линия) ───────────────
    sep = doc.add_paragraph()
    sep.paragraph_format.space_before = Pt(0)
    sep.paragraph_format.space_after = Pt(0)
    pPr = sep._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), TEAL_HEX)
    pBdr.append(bottom)
    pPr.append(pBdr)

    doc.add_paragraph()

    # ─────────────── ЗАГОЛОВОК КП ───────────────
    title = add_paragraph(doc, "КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ", bold=True, size=20,
                           color=DARK, align=WD_ALIGN_PARAGRAPH.LEFT,
                           space_before=6, space_after=4)

    subtitle = add_paragraph(doc, data["project_name"], bold=False, size=13,
                              color=TEAL, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=12)

    # ─────────────── КОМУ ───────────────
    add_paragraph(doc, f"Кому: {data['customer_full']}", bold=False, size=10,
                  color=GRAY, space_after=2)
    add_paragraph(doc, f"Контактное лицо: {data['customer_director_name']}", bold=False, size=10,
                  color=GRAY, space_after=2)
    add_paragraph(doc, f"Дата: {data['date']}", bold=False, size=10,
                  color=GRAY, space_after=16)

    # ─────────────── ОПИСАНИЕ ЗАДАЧИ ───────────────
    add_paragraph(doc, "О задаче", bold=True, size=13, color=DARK, space_before=4, space_after=4)
    add_paragraph(doc, data["task_description"], bold=False, size=11, color=DARK, space_after=16)

    # ─────────────── СОСТАВ РАБОТ ───────────────
    add_paragraph(doc, "Состав работ", bold=True, size=13, color=DARK, space_before=4, space_after=8)

    works_tbl = doc.add_table(rows=1 + len(data["works"]) + 1, cols=3)
    works_tbl.autofit = False
    works_tbl.style = 'Table Grid'

    # Ширины
    works_tbl.columns[0].width = Cm(1.2)
    works_tbl.columns[1].width = Cm(9.8)
    works_tbl.columns[2].width = Cm(3.0)

    # Заголовок таблицы
    header_row = works_tbl.rows[0]
    headers = ["№", "Наименование работ", "Часов"]
    for i, hdr in enumerate(headers):
        cell = header_row.cells[i]
        set_cell_bg(cell, TEAL_HEX)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(hdr)
        r.bold = True
        r.font.size = Pt(10)
        r.font.name = 'Calibri'
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Строки работ
    total_hours = 0
    for idx, work in enumerate(data["works"]):
        row = works_tbl.rows[idx + 1]
        hours_val = float(work["hours"].replace(",", "."))
        total_hours += hours_val

        row.cells[0].paragraphs[0].add_run(str(idx + 1)).font.size = Pt(10)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        name_run = row.cells[1].paragraphs[0].add_run(work["name"])
        name_run.font.size = Pt(10)
        name_run.font.name = 'Calibri'

        hours_run = row.cells[2].paragraphs[0].add_run(work["hours"])
        hours_run.font.size = Pt(10)
        hours_run.font.name = 'Calibri'
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Итого часов
    last_row = works_tbl.rows[-1]
    set_cell_bg(last_row.cells[0], "F2F2F2")
    set_cell_bg(last_row.cells[1], "F2F2F2")
    set_cell_bg(last_row.cells[2], "F2F2F2")

    last_row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = last_row.cells[1].paragraphs[0].add_run("Итого:")
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = 'Calibri'

    last_row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = last_row.cells[2].paragraphs[0].add_run(f"{total_hours:,.2f}".replace(",", " ").replace(".", ","))
    r2.bold = True
    r2.font.size = Pt(10)
    r2.font.name = 'Calibri'

    doc.add_paragraph()

    # ─────────────── СТОИМОСТЬ И СРОКИ ───────────────
    info_tbl = doc.add_table(rows=1, cols=2)
    info_tbl.autofit = False
    set_no_borders(info_tbl)
    info_tbl.columns[0].width = Cm(9)
    info_tbl.columns[1].width = Cm(5)

    # Стоимость
    cost_cell = info_tbl.cell(0, 0)
    # Цветной блок
    set_cell_bg(cost_cell, "F0FAFA")
    cp = cost_cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(8)
    r = cp.add_run("💰  Стоимость работ\n")
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = 'Calibri'
    r.font.color.rgb = TEAL

    r2 = cp.add_run(f"{data['total_amount']} руб.")
    r2.bold = True
    r2.font.size = Pt(18)
    r2.font.name = 'Calibri'
    r2.font.color.rgb = DARK

    r3 = cp.add_run(f"\n(в т.ч. НДС 5% — {data['vat_amount']} руб.)")
    r3.font.size = Pt(9)
    r3.font.name = 'Calibri'
    r3.font.color.rgb = GRAY

    # Сроки
    dates_cell = info_tbl.cell(0, 1)
    set_cell_bg(dates_cell, "F0FAFA")
    dp = dates_cell.paragraphs[0]
    dp.paragraph_format.space_before = Pt(8)
    dp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    rd = dp.add_run("📅  Сроки\n")
    rd.bold = True
    rd.font.size = Pt(11)
    rd.font.name = 'Calibri'
    rd.font.color.rgb = TEAL

    rd2 = dp.add_run(f"{data['start_date']} —\n{data['end_date']}")
    rd2.font.size = Pt(11)
    rd2.font.name = 'Calibri'
    rd2.font.color.rgb = DARK
    rd2.bold = True

    doc.add_paragraph()

    # ─────────────── ПОРЯДОК ОПЛАТЫ ───────────────
    add_paragraph(doc, "Порядок оплаты", bold=True, size=13, color=DARK, space_before=8, space_after=6)

    payments = data.get("payments", [])
    pay_tbl = doc.add_table(rows=len(payments), cols=3)
    pay_tbl.autofit = False
    set_no_borders(pay_tbl)
    pay_tbl.columns[0].width = Cm(1.5)
    pay_tbl.columns[1].width = Cm(10.0)
    pay_tbl.columns[2].width = Cm(2.5)

    for i, pay in enumerate(payments):
        row = pay_tbl.rows[i]
        # Буллет
        b = row.cells[0].paragraphs[0].add_run("●")
        b.font.color.rgb = TEAL
        b.font.size = Pt(10)
        b.font.name = 'Calibri'

        d = row.cells[1].paragraphs[0].add_run(pay["desc"])
        d.font.size = Pt(10)
        d.font.name = 'Calibri'
        d.font.color.rgb = DARK

        a = row.cells[2].paragraphs[0].add_run(pay["amount"])
        a.font.size = Pt(10)
        a.font.name = 'Calibri'
        a.font.color.rgb = DARK
        a.bold = True
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()

    # ─────────────── ГАРАНТИЯ ───────────────
    guar_p = add_paragraph(doc, "", space_before=4, space_after=16)
    g1 = guar_p.add_run("🛡  Гарантия: ")
    g1.bold = True
    g1.font.size = Pt(11)
    g1.font.color.rgb = TEAL
    g1.font.name = 'Calibri'
    g2 = guar_p.add_run(data.get("warranty_text", "3 месяца на выполненные работы"))
    g2.font.size = Pt(11)
    g2.font.color.rgb = DARK
    g2.font.name = 'Calibri'

    # ─────────────── РАЗДЕЛИТЕЛЬ ───────────────
    sep2 = doc.add_paragraph()
    sep2.paragraph_format.space_before = Pt(0)
    sep2.paragraph_format.space_after = Pt(8)
    pPr2 = sep2._p.get_or_add_pPr()
    pBdr2 = OxmlElement('w:pBdr')
    bot2 = OxmlElement('w:bottom')
    bot2.set(qn('w:val'), 'single')
    bot2.set(qn('w:sz'), '4')
    bot2.set(qn('w:space'), '1')
    bot2.set(qn('w:color'), TEAL_HEX)
    pBdr2.append(bot2)
    pPr2.append(pBdr2)

    # ─────────────── ФУТЕР: Контакты ───────────────
    footer_tbl = doc.add_table(rows=1, cols=2)
    footer_tbl.autofit = False
    set_no_borders(footer_tbl)
    footer_tbl.columns[0].width = Cm(9)
    footer_tbl.columns[1].width = Cm(5)

    fc1 = footer_tbl.cell(0, 0)
    fp1 = fc1.paragraphs[0]
    r = fp1.add_run("ООО «Гравити Групп»\n")
    r.bold = True
    r.font.size = Pt(10)
    r.font.name = 'Calibri'
    r.font.color.rgb = DARK

    r2 = fp1.add_run(f"Директор: {data.get('contractor_director_name', 'Яборов А.В.')}")
    r2.font.size = Pt(9)
    r2.font.name = 'Calibri'
    r2.font.color.rgb = GRAY

    fc2 = footer_tbl.cell(0, 1)
    fp2 = fc2.paragraphs[0]
    fp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r3 = fp2.add_run("По всем вопросам:\n")
    r3.font.size = Pt(9)
    r3.font.name = 'Calibri'
    r3.font.color.rgb = GRAY

    r4 = fp2.add_run(data.get("contractor_contact", ""))
    r4.font.size = Pt(9)
    r4.font.name = 'Calibri'
    r4.font.color.rgb = TEAL

    doc.save(output_path)
    print(f"КП сохранено: {output_path}")


# ─────────────── ДАННЫЕ ───────────────
if __name__ == "__main__":
    data = {
        "project_name": "Замена интеграции Dashly → SendPulse в приложении IIM",
        "date": "17 февраля 2026 г.",

        "customer_full": "ООО «ПЛАТФОРМА» (ИНН 5902054490)",
        "customer_director_name": "Гуляев Евгений Артурович",

        "task_description": (
            "Требуется заменить интеграцию с сервисом Dashly на сервис SendPulse "
            "в мобильном приложении IIM (Israel Inspires Me). "
            "Работы включают: настройку SDK SendPulse, перенос логики "
            "рассылок и push-уведомлений, тестирование корректности "
            "отправки событий и обратную совместимость с текущей архитектурой."
        ),

        "works": [
            {"name": "Разработка Back-End (Общий)", "hours": "24,00"},
            {"name": "Тестирование", "hours": "7,20"},
            {"name": "Ведение проекта", "hours": "3,60"},
        ],

        "total_amount": "126 240",
        "vat_amount": "6 011,43",

        "start_date": "17 февраля 2026 г.",
        "end_date": "3 марта 2026 г.",

        "payments": [
            {"desc": "Аванс — 30% от стоимости, в течение 2 рабочих дней с момента подписания", "amount": "37 872 руб."},
            {"desc": "37% от стоимости — в течение 30 рабочих дней, но не позже подписания Акта", "amount": "46 708 руб."},
            {"desc": "Оставшаяся часть — в течение 5 рабочих дней после подписания Акта", "amount": "41 660 руб."},
        ],

        "warranty_text": "3 (три) месяца на все выполненные работы",

        "contractor_director_name": "Яборов Андрей Владимирович",
        "contractor_contact": "gravity-group.ru",
    }

    output = "/home/clawdbot/.openclaw/workspace/drafts/KP_IIM_Gravity_2026-02-17.docx"
    generate_kp(data, output)
