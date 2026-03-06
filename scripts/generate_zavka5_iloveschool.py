#!/usr/bin/env python3
"""
Генератор Заявки №5 для ILoveSchool (ООО "ПЛАТФОРМА")
Данные из: файл Excel (оценка заявок №4 и №5)
Шаблон: file_81 (Инплайн Заявка №4)
"""

from docx import Document
from lxml import etree
from docx.oxml.ns import qn

TEMPLATE_PATH = '/home/clawdbot/.openclaw/media/inbound/file_81---2141837f-9e85-45e3-9135-f0ffbf7ac458.docx'
OUTPUT_PATH = '/home/clawdbot/.openclaw/workspace/docs/zavka5_ILoveSchool.docx'

ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def replace_in_paragraph(paragraph, old_text, new_text):
    """Заменить текст, разбитый на run'ы."""
    pieces = []
    for run_elem in paragraph.findall('.//w:r', ns):
        for t in run_elem.findall('.//w:t', ns):
            pieces.append(t)
    if not pieces:
        return False
    full = ''.join(t.text or '' for t in pieces)
    if old_text not in full:
        return False
    new_full = full.replace(old_text, new_text, 1)
    pieces[0].text = new_full
    pieces[0].set(qn('xml:space'), 'preserve')
    for t in pieces[1:]:
        t.text = ''
    return True


def replace_in_all_paragraphs(doc, replacements):
    """Заменить текст во всех параграфах документа (включая таблицы)."""
    body = doc.element.body
    all_paras = list(body.iter(qn('w:p')))
    results = {}
    for old_text, new_text in replacements:
        count = 0
        for para in all_paras:
            if replace_in_paragraph(para, old_text, new_text):
                count += 1
        results[old_text] = count
    return results


def replace_sdt_content(body, alias, new_text):
    """Заменить содержимое ВСЕХ SDT с данным alias."""
    found = False
    for sdt in body.findall('.//w:sdt', ns):
        sdtPr = sdt.find('.//w:sdtPr', ns)
        if sdtPr is not None:
            a = sdtPr.find('.//w:alias', ns)
            if a is not None and a.get(qn('w:val')) == alias:
                for t in sdt.findall('.//w:t', ns):
                    t.text = new_text
                found = True
    return found


def update_property(doc, prop_type, prop_name, new_value):
    """Обновить свойство документа."""
    for rel in doc.part.package.rels.values():
        reltype = str(rel.reltype).lower()
        if prop_type == 'app' and ('extended' in reltype or 'app' in reltype):
            ns_uri = 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'
        elif prop_type == 'core' and ('core-properties' in reltype or 'metadata/core' in reltype):
            ns_uri = 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties'
        else:
            continue
        try:
            root = etree.fromstring(rel.target_part.blob)
            elem = root.find(f'{{{ns_uri}}}{prop_name}')
            if elem is not None:
                elem.text = new_value
                rel.target_part._blob = etree.tostring(
                    root, xml_declaration=True, encoding='UTF-8', standalone=True)
                return True
        except:
            pass
    return False


def delete_table_row(table, row_index):
    """Удалить строку таблицы по индексу."""
    tbl = table._tbl
    rows = tbl.findall(qn('w:tr'))
    if row_index < len(rows):
        tbl.remove(rows[row_index])
        return True
    return False


def replace_cell_text(cell, new_text):
    """Заменить текст в ячейке таблицы (первый параграф)."""
    for para in cell.paragraphs:
        pieces = []
        for run_elem in para._p.findall('.//w:r', ns):
            for t in run_elem.findall('.//w:t', ns):
                pieces.append(t)
        if pieces:
            pieces[0].text = new_text
            pieces[0].set(qn('xml:space'), 'preserve')
            for t in pieces[1:]:
                t.text = ''
            return True
    return False


def main():
    print("Загружаем шаблон...")
    doc = Document(TEMPLATE_PATH)
    body = doc.element.body

    # ================================================================
    # 1. SDT — Реквизиты заказчика
    # ================================================================
    CUSTOMER_NAME = 'Общество с ограниченной ответственностью "ПЛАТФОРМА"'
    CUSTOMER_INN = '5902054490'

    replace_sdt_content(body, "Наименование Заказчика", CUSTOMER_NAME)
    replace_sdt_content(body, "ИНН Заказчика", CUSTOMER_INN)
    update_property(doc, 'app', 'Company', CUSTOMER_NAME)
    update_property(doc, 'core', 'contentStatus', CUSTOMER_INN)

    # ================================================================
    # 2. Замены в параграфах (все параграфы включая таблицы)
    # ================================================================
    replacements = [
        # Номер и договор
        ("ЗАЯВКА №4",                              "ЗАЯВКА №5"),
        ("к Договору 14/25",                       "к Договору 66/2023"),
        ("«16» мая 2025",                          "«13» ноября 2023"),
        ("по Договору №14/25",                     "по Договору №66/2023"),
        ("к Заявке №4 от",                         "к Заявке №5 от"),

        # Даты заявки и работ
        ("«05» февраля 2026",                      "«17» февраля 2026"),
        ("«5» февраля 2026",                       "«17» февраля 2026"),
        ("«27» февраля 2026",                      "«3» марта 2026"),

        # Директор заказчика (таблица сторон)
        ("Шевшелева Сергея Викторовича",           "Гуляева Евгения Артуровича"),

        # Подпись заказчика (таблица подписей)
        ("Генеральный директор ООО «Инплайн»",    "Генеральный директор ООО «ПЛАТФОРМА»"),
        ("Шевшелев Сергей Викторович",             "Гуляев Евгений Артурович"),
        ("С.В. Шевшелев",                          "Е.А. Гуляев"),

        # Итоговая сумма и НДС
        # "166 057 (... рублей, 50 копеек" → "126 240 (... рублей, 00 копеек"
        (
            "166 057 (Сто шестьдесят шесть тысяч пятьдесят семь) рублей, 50 копеек",
            "126 240 (Сто двадцать шесть тысяч двести сорок) рублей, 00 копеек"
        ),
        ("7\xa0907,50",                            "6\xa0011,43"),
        ("7 907,50",                               "6 011,43"),

        # Первый платёж (30%)
        (
            "49 817 (Сорок девять тысяч восемьсот семнадцать) рублей 25 копеек",
            "37 872 (Тридцать семь тысяч восемьсот семьдесят два) рубля 00 копеек"
        ),

        # Второй платёж (37%)
        (
            "61 441 (Шестьдесят одна тысяча четыреста сорок один) рубль 28 копеек",
            "46 708 (Сорок шесть тысяч семьсот восемь) рублей 80 копеек"
        ),
    ]

    results = replace_in_all_paragraphs(doc, replacements)
    for old, count in results.items():
        status = "✅" if count > 0 else "⚠️ НЕ НАЙДЕНО"
        print(f"  {status} [{count}x] {old[:60]}")

    # ================================================================
    # 3. TABLE 2 — Работы
    #    Шаблон: Дизайнер(8), Веб разработчик(23), Специалист по тестированию(10),
    #             Руководитель проекта(7), Аналитик(1)
    #    Заявка 5 ILS: Разработка Back-End(24), Тестирование(7,20), Ведение проекта(3,60)
    # ================================================================
    table2 = doc.tables[1]  # 0-indexed

    # Row 1: label строка "Разработка мобильной версии..." → "Общая стоимость проекта"
    row1_cell1 = table2.rows[1].cells[1]
    replace_cell_text(row1_cell1, "Общая стоимость проекта (Предварительная оценка)")

    # Row 2: Дизайнер → Разработка Back-End (Общий) + часы 8,00 → 24,00
    replace_cell_text(table2.rows[2].cells[1], "Разработка Back-End (Общий)")
    replace_cell_text(table2.rows[2].cells[2], "24,00")

    # Row 3: Веб разработчик → Тестирование + 23,00 → 7,20
    replace_cell_text(table2.rows[3].cells[1], "Тестирование")
    replace_cell_text(table2.rows[3].cells[2], "7,20")

    # Row 4: Специалист по тестированию → Ведение проекта + 10,00 → 3,60
    replace_cell_text(table2.rows[4].cells[1], "Ведение проекта")
    replace_cell_text(table2.rows[4].cells[2], "3,60")

    # Удаляем строки 5 и 6 (Руководитель проекта, Аналитик)
    # Удаляем начиная с конца, чтобы не сбить индексы
    delete_table_row(table2, 6)  # Аналитик
    delete_table_row(table2, 5)  # Руководитель проекта

    print(f"\n✅ Таблица работ обновлена: {len(table2.rows)} строк")

    # ================================================================
    # 4. TABLE 4 — ТЗ
    #    Одна задача: Заменить интеграцию с Dashly на SendPulse (IIM)
    # ================================================================
    table4 = doc.tables[3]

    # Row 0: обновляем задачу и описание
    replace_cell_text(
        table4.rows[0].cells[0],
        "Заменить интеграцию с Dashly на сервис SendPulse (IIM)"
    )
    replace_cell_text(
        table4.rows[0].cells[1],
        "Требуется заменить интеграцию с сервисом Dashly на сервис SendPulse "
        "в приложении IIM (Israel Inspires Me). "
        "Включает: настройку SDK SendPulse, перенос логики рассылок/уведомлений, "
        "тестирование корректности отправки событий."
    )

    # Удаляем лишние строки (3, 2, 1 — с конца)
    for idx in [3, 2, 1]:
        delete_table_row(table4, idx)

    print(f"✅ Таблица ТЗ обновлена: {len(table4.rows)} строк")

    # ================================================================
    # 5. Сохраняем
    # ================================================================
    doc.save(OUTPUT_PATH)
    print(f"\n✅ Заявка №5 ILoveSchool → {OUTPUT_PATH}")

    # Верификация
    doc2 = Document(OUTPUT_PATH)
    body2 = doc2.element.body
    sdts = body2.findall('.//w:sdt', ns)
    numPrs = body2.findall('.//w:numPr', ns)
    print(f"📊 SDT: {len(sdts)} | numPr: {len(numPrs)} | Tables: {len(doc2.tables)}")
    for i, t in enumerate(doc2.tables):
        print(f"  Table {i+1}: {len(t.rows)} rows")


if __name__ == "__main__":
    main()
