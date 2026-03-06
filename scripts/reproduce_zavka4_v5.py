#!/usr/bin/env python3
"""
Заявка №4 v5 — ПОЛНОЕ СООТВЕТСТВИЕ (все 8 проблем исправлены)
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PURPLE = '5F497A'
NBSP = '\xa0'  # Неразрывный пробел

def set_table_borders(table, borders_config):
    """
    Установка границ таблицы с точной конфигурацией.
    borders_config = dict: border_name -> {'val': ..., 'color': ..., 'sz': ...} или 'nil'
    """
    tbl = table._element
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    
    tblBorders = OxmlElement('w:tblBorders')
    for name, cfg in borders_config.items():
        border = OxmlElement(f'w:{name}')
        if cfg == 'nil':
            border.set(qn('w:val'), 'nil')
        else:
            border.set(qn('w:val'), cfg.get('val', 'single'))
            border.set(qn('w:sz'), cfg.get('sz', '4'))
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), cfg.get('color', '000000'))
        tblBorders.append(border)
    tblPr.append(tblBorders)

def set_cell_borders(cell, borders_config):
    """Установка границ на уровне ячейки."""
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    old = tcPr.find(qn('w:tcBorders'))
    if old is not None:
        tcPr.remove(old)
    
    tcBorders = OxmlElement('w:tcBorders')
    for name, cfg in borders_config.items():
        border = OxmlElement(f'w:{name}')
        border.set(qn('w:val'), cfg.get('val', 'single'))
        border.set(qn('w:sz'), cfg.get('sz', '4'))
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), cfg.get('color', '000000'))
        tcBorders.append(border)
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    """Установка фонового цвета ячейки."""
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    # Удалить старый shading
    old = tcPr.find(qn('w:shd'))
    if old is not None:
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color_hex)
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)

def set_cell_width(cell, width_cm):
    """Установить ширину ячейки."""
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(int(width_cm * 567)))  # cm to twips
    tcW.set(qn('w:type'), 'dxa')
    old = tcPr.find(qn('w:tcW'))
    if old is not None:
        tcPr.remove(old)
    tcPr.append(tcW)

# === СОЗДАНИЕ ДОКУМЕНТА ===
doc = Document()

# Секция
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(3.0)
section.right_margin = Cm(1.5)
section.header_distance = Emu(449580)  # Точное значение из оригинала
section.footer_distance = Emu(449580)

# === ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ ===
def add_para(text, bold=False, with_indent=False, alignment=None):
    para = doc.add_paragraph()
    if with_indent:
        para.paragraph_format.left_indent = Pt(28.35)
        para.paragraph_format.first_line_indent = Pt(-28.35)
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.2
    para.style = 'Normal'
    if alignment is not None:
        para.alignment = alignment
    run = para.add_run(text)
    run.font.name = 'Calibri'
    run.font.bold = bold
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return para

def cell_para(cell, text, bold=False, font_color=RGBColor(0,0,0), alignment=None):
    """Настроить первый параграф ячейки."""
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.2
    if alignment is not None:
        para.alignment = alignment
    run = para.add_run(text)
    run.font.name = 'Calibri'
    run.font.bold = bold
    run.font.size = Pt(12)
    run.font.color.rgb = font_color
    return run

# ==============================================================
# ЗАГОЛОВОК
# ==============================================================
add_para("ЗАЯВКА №4 к Договору 14/25 от «16» мая 2025 г.", bold=True, with_indent=True)
add_para("Дата Заявки: «05» февраля 2026 г.", bold=True, with_indent=True)

# ==============================================================
# ТАБЛИЦА 1: СТОРОНЫ
# ==============================================================
table1 = doc.add_table(rows=2, cols=2)
table1.style = 'Normal Table'
set_table_borders(table1, {
    'top': 'nil', 'left': 'nil', 'bottom': 'nil', 'right': 'nil',
    'insideH': {'val': 'single', 'color': PURPLE, 'sz': '4'},
    'insideV': {'val': 'single', 'color': PURPLE, 'sz': '4'},
})
table1.columns[0].width = Cm(4.25)
table1.columns[1].width = Cm(12.26)

cell_para(table1.rows[0].cells[0], "Заказчик", bold=True)
cell_para(table1.rows[0].cells[1],
    "ИНН \nв лице директора Шевшелева Сергея Викторовича\nдействующего на основании Устава")
cell_para(table1.rows[1].cells[0], "Исполнитель", bold=True)
cell_para(table1.rows[1].cells[1],
    "Общество с ограниченной ответственностью «Гравити Групп»\n"
    "ИНН 5908999996\n"
    "в лице Директора Яборова Андрея Владимировича, действующего на основании Устава")

# ==============================================================
# "согласовали..." → ТАБЛИЦА РАБОТ → текст стоимости
# (ИСПРАВЛЕНИЕ #1: правильный порядок!)
# ==============================================================
add_para("согласовали следующий условия выполнения работ по Заявке к Договору:")

# ТАБЛИЦА 2: РАБОТЫ (сразу после "согласовали...")
table2 = doc.add_table(rows=7, cols=3)
table2.style = 'Normal Table'
# ИСПРАВЛЕНИЕ #2: ТОЛЬКО insideH, без insideV
set_table_borders(table2, {
    'insideH': {'val': 'single', 'color': PURPLE, 'sz': '4'},
})
table2.columns[0].width = Cm(1.25)
table2.columns[1].width = Cm(11.00)
table2.columns[2].width = Cm(3.86)

# Заголовок таблицы работ
headers = ['№', 'Наименование и описание работ', 'Оценка трудоёмкости в часах']
for i, header in enumerate(headers):
    cell = table2.rows[0].cells[i]
    para = cell.paragraphs[0]
    para.paragraph_format.left_indent = Pt(5.40)
    para.paragraph_format.first_line_indent = Pt(-5.40)
    para.paragraph_format.line_spacing = 1.2
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(header)
    run.font.name = 'Calibri'
    run.font.bold = True
    run.font.size = Pt(12)

# Категория
cell_para(table2.rows[1].cells[1],
    "Разработка мобильной версии приложения (Приложение №1 к Заявке)", bold=True)

# Специалисты
specialists = [
    ("Дизайнер", "8,00"),
    ("Веб разработчик", "23,00"),
    ("Специалист по тестированию", "10,00"),
    ("Руководитель проекта", "7,00"),
    ("Аналитик", "1,00"),
]
for idx, (name, hours) in enumerate(specialists, start=2):
    para = table2.rows[idx].cells[1].paragraphs[0]
    para.paragraph_format.line_spacing = 1.2
    run = para.add_run(name)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    
    para = table2.rows[idx].cells[2].paragraphs[0]
    para.paragraph_format.line_spacing = 1.2
    run = para.add_run(hours)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)

# ТЕКСТОВЫЕ ПАРАГРАФЫ (после таблицы работ!)
# ИСПРАВЛЕНИЕ #6: неразрывные пробелы
add_para(f"Предварительная стоимость работ по Заявке составляет 166{NBSP}057 (Сто шестьдесят шесть тысяч пятьдесят семь) рублей, 50 копеек, в т.ч. НДС 5% - 7{NBSP}907,50 рублей.", with_indent=True)
add_para("Сроки выполнения Заявки:", with_indent=True)
add_para("Дата начала выполнения работ: «5» февраля 2026 г.", with_indent=True)
add_para("Дата окончания выполнения работ: «27» февраля 2026 г.", with_indent=True)
add_para("Порядок оплаты работ по Заявке:", with_indent=True)
add_para(f"30% планируемой стоимости работ в размере 49{NBSP}817 (Сорок девять тысяч восемьсот семнадцать) рублей 25 копеек оплачиваются в течение 2 (Двух) рабочих с момента подписания Заявки.", with_indent=True)
add_para(f"37% планируемой стоимости работ в размере 61{NBSP}441 (Шестьдесят одна тысяча четыреста сорок один) рубль 28 копеек оплачиваются в течение 30 (Тридцати) рабочих с момента подписания Заявки, но не позже даты подписания Акта по Заявке.", with_indent=True)
add_para("Оплата оставшейся части общей стоимости работ, рассчитанной на основании фактических трудозатрат Исполнителя, и указанных в Акте по производится в течение 5 (Пяти) рабочих дней со дня подписания Акта по Заявке. По согласованию с Исполнителем оплата оставшейся части общей стоимости выполненных работ может быть отсрочена до 20 (Двадцати) рабочих дней после подписания Акта сдачи-приемки работ.", with_indent=True)
add_para("Срок гарантии на выполненные работы Исполнителем составляет 3 (Три) месяца.", with_indent=True)
add_para("Если для выполнения Заявки потребуются проработка предметной области, разработка 3D моделей, то для выполнения работ со стороны Исполнителя привлекается аналитик, разработчик, трудозатраты будут отражены в Акте по Заявке.", with_indent=True)
add_para("Положения Заявки имеют приоритет над условиями Договора. В остальном, что не предусмотрено Заявкой, Стороны руководствуются условиями Договора.", with_indent=True)

# ==============================================================
# ТАБЛИЦА 3: ПОДПИСИ
# ИСПРАВЛЕНИЕ #3: только insideV, nil для остальных
# ==============================================================
table3 = doc.add_table(rows=4, cols=2)
table3.style = 'Normal Table'
set_table_borders(table3, {
    'top': 'nil', 'left': 'nil', 'bottom': 'nil', 'right': 'nil',
    'insideH': 'nil',
    'insideV': {'val': 'single', 'color': PURPLE, 'sz': '4'},
})
table3.columns[0].width = Cm(8.24)
table3.columns[1].width = Cm(8.24)

# Заголовки подписей
for col_idx, title in enumerate(["ЗАКАЗЧИК", "ИСПОЛНИТЕЛЬ"]):
    cell = table3.rows[0].cells[col_idx]
    set_cell_shading(cell, PURPLE)
    cell_para(cell, title, bold=True, font_color=RGBColor(255,255,255),
              alignment=WD_ALIGN_PARAGRAPH.CENTER)

# Организации
cell_para(table3.rows[1].cells[1],
    "Общество с ограниченной ответственностью «Гравити Групп»\n")

# Должности
cell_para(table3.rows[2].cells[0],
    "Генеральный директор ООО «Инплайн»\nШевшелев Сергей Викторович\n\n(на основании Устава)")
cell_para(table3.rows[2].cells[1],
    "Директор ООО «ГРАВИТИ ГРУПП»\nЯборов Андрей Владимирович\n(на основании Устава)")

# Подписи
cell_para(table3.rows[3].cells[0], "\n_____________________ / С.В. Шевшелев")
cell_para(table3.rows[3].cells[1], "\n_________________________ / А.В. Яборов")

# ==============================================================
# ИСПРАВЛЕНИЕ #5: 16 ПУСТЫХ ПАРАГРАФОВ (разрыв страницы)
# ==============================================================
for _ in range(16):
    doc.add_paragraph()

# ==============================================================
# ПРИЛОЖЕНИЕ №1
# ==============================================================
add_para("Приложение №1", bold=True, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
add_para("к Заявке №4 от «05» февраля 2026 г.", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
add_para("по Договору №14/25 от «16» мая 2025 г.", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
doc.add_paragraph()  # Пустой
add_para("Техническое задание", alignment=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()  # Пустой

# ==============================================================
# ТАБЛИЦА 4: ТЗ
# ИСПРАВЛЕНИЕ #4: cell-level borders, другие цвета/ширины
# ==============================================================
table4 = doc.add_table(rows=4, cols=2)
table4.style = 'Normal Table'
# НЕТ tblBorders!

table4.columns[0].width = Cm(5.17)
table4.columns[1].width = Cm(11.57)

GREY_TEXT = RGBColor(0x43, 0x43, 0x43)

tz_data = [
    ("Добавление новых типов банок и моделей",
     "Требуется добавить 3 вида новых банок и конвертировать модели в поддерживаемый формат: \n1. Банки с горлом 38мм\n1.1 100 мл\n1.2 150 мл\n1.3 155 мл\n"),
    ("Добавление предпросмотра капсул",
     "Требуется добавить предпросмотр капсул в этапе выбора МЖК/ТЖК/Таблеток\n1. Реализовать функционал создания изображения из модели капсулы\n2. Отобразить изображение для каждой капсулы"),
    ("Обновление логики расчетов для физики банок",
     "Требуется добавить поддержку логики расчета для новых размеров банок\n1. Требуется подогнать физические ограничения и максимальный уровень насыпки под каждый размер новых моделей"),
    ("Мобильная версия",
     "Требуется разработать дизайн мобильной версии и интегрировать его в конфигуратор"),
]

for row_idx, (title, desc) in enumerate(tz_data):
    cell0 = table4.rows[row_idx].cells[0]
    cell1 = table4.rows[row_idx].cells[1]
    
    # Shading
    set_cell_shading(cell0, 'FFFFFF')
    set_cell_shading(cell1, 'FFFFFF')
    
    # Cell-level borders
    outer = {'val': 'single', 'color': '000000', 'sz': '6'}
    inner = {'val': 'single', 'color': 'CCCCCC', 'sz': '6'}
    
    # Cell 0 borders
    borders0 = {'top': outer, 'left': outer, 'bottom': outer, 'right': outer}
    if row_idx > 0:
        borders0['top'] = inner
    
    # Cell 1 borders
    borders1 = {'top': outer, 'left': inner, 'bottom': outer, 'right': outer}
    if row_idx > 0:
        borders1['top'] = inner
    
    set_cell_borders(cell0, borders0)
    set_cell_borders(cell1, borders1)
    
    # Текст
    cell_para(cell0, title, bold=True, font_color=GREY_TEXT)
    cell_para(cell1, desc, font_color=GREY_TEXT)

# Финальный пустой параграф
doc.add_paragraph()

# === СОХРАНЕНИЕ ===
output = "/home/clawdbot/.openclaw/workspace/docs/zavka4_v5_AUDITED.docx"
doc.save(output)

print(f"✅ v5 создан: {output}")
print("\n📋 ВСЕ 8 ИСПРАВЛЕНИЙ:")
print("  1. ✅ Порядок: TABLE работы ПЕРЕД текстом стоимости")
print("  2. ✅ TABLE работы: только insideH (без insideV)")
print("  3. ✅ TABLE подписей: только insideV, nil для остальных")
print("  4. ✅ TABLE ТЗ: cell-level borders, 434343, FFFFFF, 5.17x11.57")
print("  5. ✅ 16 пустых параграфов перед Приложением")
print("  6. ✅ Неразрывные пробелы в суммах")
print("  7. ✅ Явные nil-границы в таблице сторон")
print("  8. ✅ Header/footer distance = 449580")
