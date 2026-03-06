#!/usr/bin/env python3
"""
Генератор ЗАЯВОК для Gravity Group v4 (ПОЛНОЕ СООТВЕТСТВИЕ оригиналу)
Все детали выверены по анализу 6 оригинальных документов

Использование: python3 gravity_zavka_template_v4.py
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from datetime import datetime

def set_table_borders_colored(table, color_hex='5F497A', inside_only=False):
    """
    Установка цветных границ таблицы (как в оригинале Gravity Group)
    
    Args:
        table: Таблица docx
        color_hex: Цвет границ в hex (по умолчанию 5F497A - фиолетовый)
        inside_only: Если True, красить только внутренние границы
    """
    tbl = table._element
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    
    # Создаём элемент границ
    tblBorders = OxmlElement('w:tblBorders')
    
    # Определяем какие границы красить
    if inside_only:
        borders_to_set = ['insideH', 'insideV']
    else:
        borders_to_set = ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']
    
    for border_name in borders_to_set:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')  # Толщина
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color_hex)
        tblBorders.append(border)
    
    # Удаляем старые границы если есть
    old_borders = tblPr.find(qn('w:tblBorders'))
    if old_borders is not None:
        tblPr.remove(old_borders)
    
    tblPr.append(tblBorders)

def set_cell_shading(cell, color_hex):
    """Установка фонового цвета ячейки"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color_hex)
    cell._element.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=0, bottom=0, left=108, right=108):
    """
    Установка отступов внутри ячейки (в twips, 1 pt = 20 twips)
    По умолчанию: left=right=108 twips (5.4pt) как в оригинале
    """
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    
    for margin_name, value in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if value is not None:
            margin = OxmlElement(f'w:{margin_name}')
            margin.set(qn('w:w'), str(value))
            margin.set(qn('w:type'), 'dxa')
            tcMar.append(margin)
    
    tcPr.append(tcMar)

def number_to_russian_words(num):
    """Конвертация числа в слова на русском языке"""
    if num == 0:
        return "ноль"
    
    ones = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    teens = ["десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать", 
             "пятнадцать", "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать"]
    tens = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", 
            "шестьдесят", "семьдесят", "восемьдесят", "девяносто"]
    hundreds = ["", "сто", "двести", "триста", "четыреста", "пятьсот", 
                "шестьсот", "семьсот", "восемьсот", "девятьсот"]
    
    thousands = ["", "одна", "две", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    
    def convert_hundreds(n):
        if n == 0:
            return ""
        elif n < 10:
            return ones[n]
        elif n < 20:
            return teens[n - 10]
        elif n < 100:
            return tens[n // 10] + (" " + ones[n % 10] if n % 10 != 0 else "")
        else:
            h = hundreds[n // 100]
            remainder = n % 100
            if remainder == 0:
                return h
            elif remainder < 10:
                return h + " " + ones[remainder]
            elif remainder < 20:
                return h + " " + teens[remainder - 10]
            else:
                return h + " " + tens[remainder // 10] + (" " + ones[remainder % 10] if remainder % 10 != 0 else "")
    
    def get_thousand_form(n):
        if n % 10 == 1 and n % 100 != 11:
            return "тысяча"
        elif n % 10 in [2, 3, 4] and n % 100 not in [12, 13, 14]:
            return "тысячи"
        else:
            return "тысяч"
    
    num = int(num)
    if num == 0:
        return "ноль"
    
    result = []
    
    # Миллионы
    millions = num // 1000000
    if millions > 0:
        result.append(convert_hundreds(millions))
        if millions % 10 == 1 and millions % 100 != 11:
            result.append("миллион")
        elif millions % 10 in [2, 3, 4] and millions % 100 not in [12, 13, 14]:
            result.append("миллиона")
        else:
            result.append("миллионов")
    
    # Тысячи
    thousands_num = (num % 1000000) // 1000
    if thousands_num > 0:
        if thousands_num < 10:
            result.append(thousands[thousands_num])
        elif thousands_num < 20:
            result.append(teens[thousands_num - 10])
        elif thousands_num < 100:
            t = tens[thousands_num // 10]
            o = thousands[thousands_num % 10] if thousands_num % 10 != 0 else ""
            result.append(t + (" " + o if o else ""))
        else:
            h = hundreds[thousands_num // 100]
            remainder = thousands_num % 100
            if remainder == 0:
                result.append(h)
            elif remainder < 10:
                result.append(h + " " + thousands[remainder])
            elif remainder < 20:
                result.append(h + " " + teens[remainder - 10])
            else:
                t = tens[remainder // 10]
                o = thousands[remainder % 10] if remainder % 10 != 0 else ""
                result.append(h + " " + t + (" " + o if o else ""))
        
        result.append(get_thousand_form(thousands_num))
    
    # Единицы
    ones_num = num % 1000
    if ones_num > 0:
        result.append(convert_hundreds(ones_num))
    
    return " ".join(result)

def format_money(amount):
    """
    Форматирование денег в стиле оригинала: 147 450 (пробел как разделитель)
    """
    rubles = int(amount)
    kopeks = int(round((amount - rubles) * 100))
    
    # Форматируем рубли с пробелами как разделителями тысяч
    rubles_str = f"{rubles:,}".replace(',', ' ')
    
    return rubles_str, kopeks

def money_to_words(amount):
    """Конвертация суммы в слова"""
    rubles = int(amount)
    kopeks = int(round((amount - rubles) * 100))
    
    rubles_words = number_to_russian_words(rubles).capitalize()
    
    return f"{rubles_words} рублей, {kopeks:02d} копеек"

def create_zavka_document(data):
    """
    Создание ТОЧНОЙ КОПИИ заявки Gravity Group v4
    ПОЛНОЕ СООТВЕТСТВИЕ оригиналу (все детали выверены)
    """
    
    doc = Document()
    
    # === УСТАНОВКА MARGINS ===
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)    # ⚠️ НЕ 2см!
    section.right_margin = Cm(1.5)   # ⚠️ НЕ 2см!
    
    # === ЗАГОЛОВОК (С ВИСЯЧЕЙ СТРОКОЙ!) ===
    title = doc.add_paragraph()
    title.paragraph_format.left_indent = Pt(28.35)      # ⚠️ Важно!
    title.paragraph_format.first_line_indent = Pt(-28.35)  # ⚠️ Висячая строка
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.line_spacing = 1.2           # ⚠️ Межстрочный интервал
    title.style = 'Normal'
    
    run = title.add_run(f"ЗАЯВКА №{data['zavka_number']} к Договору {data['contract_number']} от «{data['contract_date']}» г.")
    run.font.name = 'Calibri'  # ⚠️ НЕ Arial!
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    # Дата заявки
    date_para = doc.add_paragraph()
    date_para.paragraph_format.left_indent = Pt(28.35)
    date_para.paragraph_format.first_line_indent = Pt(-28.35)
    date_para.paragraph_format.space_before = Pt(6)
    date_para.paragraph_format.space_after = Pt(6)
    date_para.paragraph_format.line_spacing = 1.2
    date_para.style = 'Normal'
    
    run = date_para.add_run(f"Дата Заявки: «{data['zavka_date']}» г.")
    run.font.name = 'Calibri'
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    # === ТАБЛИЦА СТОРОН ===
    table_parties = doc.add_table(rows=2, cols=2)
    table_parties.style = 'Normal Table'
    set_table_borders_colored(table_parties, color_hex='5F497A', inside_only=True)
    
    # Ширина колонок (точные значения из оригинала)
    table_parties.columns[0].width = Cm(4.25)
    table_parties.columns[1].width = Cm(12.26)
    
    # Заголовки
    cell = table_parties.rows[0].cells[0]
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.2
    para.style = 'Normal'
    run = para.add_run("Заказчик")
    run.font.name = 'Calibri'
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    cell = table_parties.rows[0].cells[1]
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.2
    para.style = 'Normal'
    customer_text = (
        f"{data['customer_name']}\n"
        f"ИНН {data['customer_inn']}\n"
        f"в лице директора {data['customer_director']}\n"
        f"действующего на основании Устава"
    )
    run = para.add_run(customer_text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    # Исполнитель
    cell = table_parties.rows[1].cells[0]
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.2
    para.style = 'Normal'
    run = para.add_run("Исполнитель")
    run.font.name = 'Calibri'
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    cell = table_parties.rows[1].cells[1]
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.2
    para.style = 'Normal'
    executor_text = (
        "Общество с ограниченной ответственностью «Гравити Групп»\n"
        "ИНН 5908999996\n"
        "в лице Директора Яборова Андрея Владимировича, действующего на основании Устава"
    )
    run = para.add_run(executor_text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    # === ОСНОВНОЙ ТЕКСТ ===
    def add_normal_paragraph(text, with_indent=False):
        """Вспомогательная функция для добавления параграфа"""
        para = doc.add_paragraph()
        if with_indent:
            para.paragraph_format.left_indent = Pt(28.35)
            para.paragraph_format.first_line_indent = Pt(-28.35)
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.line_spacing = 1.2
        para.style = 'Normal'
        run = para.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return para
    
    add_normal_paragraph("согласовали следующий условия выполнения работ по Заявке к Договору:")
    
    # Стоимость
    rubles_str, kopeks = format_money(data['cost_total'])
    cost_text = f"Предварительная стоимость работ по Заявке составляет {rubles_str} ({money_to_words(data['cost_total'])})"
    if data.get('cost_vat_percent'):
        vat_amount = data['cost_total'] * data['cost_vat_percent'] / (100 + data['cost_vat_percent'])
        vat_str, vat_kop = format_money(vat_amount)
        cost_text += f", в т.ч. НДС {data['cost_vat_percent']}% - {vat_str},{vat_kop:02d} рублей"
    cost_text += "."
    add_normal_paragraph(cost_text, with_indent=True)
    
    # Сроки
    add_normal_paragraph("Сроки выполнения Заявки:", with_indent=True)
    add_normal_paragraph(f"Дата начала выполнения работ: «{data['work_start']}» г.", with_indent=True)
    add_normal_paragraph(f"Дата окончания выполнения работ: «{data['work_end']}» г.", with_indent=True)
    
    # Порядок оплаты
    add_normal_paragraph("Порядок оплаты работ по Заявке:", with_indent=True)
    
    pay30_str, pay30_kop = format_money(data['payment_30_percent'])
    add_normal_paragraph(
        f"30% планируемой стоимости работ в размере {pay30_str} ({money_to_words(data['payment_30_percent'])}) "
        f"оплачиваются в течение 2 (Двух) рабочих с момента подписания Заявки.",
        with_indent=True
    )
    
    pay37_str, pay37_kop = format_money(data['payment_37_percent'])
    add_normal_paragraph(
        f"37% планируемой стоимости работ в размере {pay37_str} ({money_to_words(data['payment_37_percent'])}) "
        f"оплачиваются в течение 30 (Тридцати) рабочих с момента подписания Заявки, но не позже даты подписания Акта по Заявке.",
        with_indent=True
    )
    
    add_normal_paragraph(
        "Оплата оставшейся части общей стоимости работ, рассчитанной на основании фактических "
        "трудозатрат Исполнителя, и указанных в Акте по производится в течение 5 (Пяти) рабочих "
        "дней со дня подписания Акта по Заявке. По согласованию с Исполнителем оплата оставшейся "
        "части общей стоимости выполненных работ может быть отсрочена до 20 (Двадцати) рабочих "
        "дней после подписания Акта сдачи-приемки работ.",
        with_indent=True
    )
    
    # Гарантия
    warranty_word = number_to_russian_words(data['warranty_months']).capitalize()
    add_normal_paragraph(
        f"Срок гарантии на выполненные работы Исполнителем составляет {data['warranty_months']} "
        f"({warranty_word}) месяца.",
        with_indent=True
    )
    
    # Доп условие
    add_normal_paragraph(
        "Если для выполнения Заявки потребуются проработка предметной области, разработка 3D моделей, "
        "то для выполнения работ со стороны Исполнителя привлекается аналитик, разработчик, "
        "трудозатраты будут отражены в Акте по Заявке.",
        with_indent=True
    )
    
    # Приоритет
    add_normal_paragraph(
        "Положения Заявки имеют приоритет над условиями Договора. В остальном, что не предусмотрено "
        "Заявкой, Стороны руководствуются условиями Договора.",
        with_indent=True
    )
    
    # === ТАБЛИЦА РАБОТ ===
    # Подсчёт строк
    total_rows = 1  # Заголовок
    for work_cat in data['works']:
        total_rows += 1  # Категория
        total_rows += len(work_cat['items'])  # Специалисты
    
    table_works = doc.add_table(rows=total_rows, cols=3)
    table_works.style = 'Normal Table'
    set_table_borders_colored(table_works, color_hex='5F497A', inside_only=True)
    
    # Ширина колонок
    table_works.columns[0].width = Cm(1.25)
    table_works.columns[1].width = Cm(11.00)
    table_works.columns[2].width = Cm(3.86)
    
    # Заголовок
    headers = ['№', 'Наименование и описание работ', 'Оценка трудоёмкости в часах']
    for i, header in enumerate(headers):
        cell = table_works.rows[0].cells[i]
        para = cell.paragraphs[0]
        para.paragraph_format.left_indent = Pt(5.40)
        para.paragraph_format.first_line_indent = Pt(-5.40)
        para.paragraph_format.line_spacing = 1.2
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.style = 'Normal'
        run = para.add_run(header)
        run.font.name = 'Calibri'
        run.font.bold = True
        run.font.size = Pt(12)
    
    # Заполнение
    row_idx = 1
    for work_cat in data['works']:
        cell = table_works.rows[row_idx].cells[1]
        para = cell.paragraphs[0]
        para.paragraph_format.line_spacing = 1.2
        para.style = 'Normal'
        run = para.add_run(work_cat['category'])
        run.font.name = 'Calibri'
        run.font.bold = True
        run.font.size = Pt(12)
        row_idx += 1
        
        for item in work_cat['items']:
            cell = table_works.rows[row_idx].cells[1]
            para = cell.paragraphs[0]
            para.paragraph_format.line_spacing = 1.2
            para.style = 'Normal'
            run = para.add_run(item['specialist'])
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            
            cell = table_works.rows[row_idx].cells[2]
            para = cell.paragraphs[0]
            para.paragraph_format.line_spacing = 1.2
            para.style = 'Normal'
            # Формат часов: 23,00 (запятая как разделитель!)
            run = para.add_run(f"{item['hours']:.2f}".replace('.', ','))
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            
            row_idx += 1
    
    # === ТАБЛИЦА ПОДПИСЕЙ ===
    table_signs = doc.add_table(rows=4, cols=2)
    table_signs.style = 'Normal Table'
    set_table_borders_colored(table_signs, color_hex='5F497A', inside_only=True)
    
    # Ширина колонок (равные)
    table_signs.columns[0].width = Cm(8.24)
    table_signs.columns[1].width = Cm(8.24)
    
    # Заголовки (с фиолетовым фоном и белым текстом)
    cell = table_signs.rows[0].cells[0]
    set_cell_shading(cell, '5F497A')  # Фиолетовый фон
    para = cell.paragraphs[0]
    para.paragraph_format.line_spacing = 1.2
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.style = 'Normal'
    run = para.add_run("ЗАКАЗЧИК")
    run.font.name = 'Calibri'
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(255, 255, 255)  # Белый текст!
    
    cell = table_signs.rows[0].cells[1]
    set_cell_shading(cell, '5F497A')
    para = cell.paragraphs[0]
    para.paragraph_format.line_spacing = 1.2
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.style = 'Normal'
    run = para.add_run("ИСПОЛНИТЕЛЬ")
    run.font.name = 'Calibri'
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(255, 255, 255)
    
    # Организации
    cell = table_signs.rows[1].cells[0]
    para = cell.paragraphs[0]
    para.paragraph_format.line_spacing = 1.2
    para.style = 'Normal'
    
    cell = table_signs.rows[1].cells[1]
    para = cell.paragraphs[0]
    para.paragraph_format.line_spacing = 1.2
    para.style = 'Normal'
    run = para.add_run("Общество с ограниченной ответственностью «Гравити Групп»")
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    
    # Должности
    cell = table_signs.rows[2].cells[0]
    para = cell.paragraphs[0]
    para.paragraph_format.line_spacing = 1.2
    para.style = 'Normal'
    run = para.add_run(
        f"Генеральный директор {data['customer_name']}\n"
        f"{data['customer_director']}\n\n"
        f"(на основании Устава)"
    )
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    
    cell = table_signs.rows[2].cells[1]
    para = cell.paragraphs[0]
    para.paragraph_format.line_spacing = 1.2
    para.style = 'Normal'
    run = para.add_run(
        "Директор ООО «ГРАВИТИ ГРУПП»\n"
        "Яборов Андрей Владимирович\n"
        "(на основании Устава)"
    )
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    
    # Подписи
    cell = table_signs.rows[3].cells[0]
    para = cell.paragraphs[0]
    para.paragraph_format.line_spacing = 1.2
    para.style = 'Normal'
    run = para.add_run(f"_____________________ / {data['customer_director_short']}")
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    
    cell = table_signs.rows[3].cells[1]
    para = cell.paragraphs[0]
    para.paragraph_format.line_spacing = 1.2
    para.style = 'Normal'
    run = para.add_run("_________________________ / А.В. Яборов")
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    
    return doc

# === ПРИМЕР ===
if __name__ == "__main__":
    example_data = {
        'zavka_number': '4',
        'contract_number': '14/25',
        'contract_date': '16» мая 2025',  # Обратите внимание на формат!
        'zavka_date': '05» февраля 2026',
        'cost_total': 166057.50,
        'cost_vat_percent': 5,
        'work_start': '5» февраля 2026',
        'work_end': '27» февраля 2026',
        'payment_30_percent': 49817.25,
        'payment_37_percent': 61441.28,
        'warranty_months': 3,
        'customer_name': 'ООО "Инплайн"',
        'customer_inn': '5902054490',
        'customer_director': 'Шевшелев Сергей Викторович',
        'customer_director_short': 'С.В. Шевшелев',
        'works': [
            {'category': 'Разработка мобильной версии приложения (Приложение №1 к Заявке)', 'items': [
                {'specialist': 'Дизайнер', 'hours': 8.0},
                {'specialist': 'Веб разработчик', 'hours': 23.0},
                {'specialist': 'Специалист по тестированию', 'hours': 10.0}
            ]}
        ]
    }
    
    doc = create_zavka_document(example_data)
    
    output_path = "/home/clawdbot/.openclaw/workspace/docs/gravity_zavka_v4_PERFECT.docx"
    doc.save(output_path)
    
    print(f"✅ v4 с ПОЛНЫМ соответствием оригиналу создан: {output_path}")
    print("\n📋 v4 ВКЛЮЧАЕТ ВСЕ ДЕТАЛИ:")
    print("  ✅ Margins: 3см слева, 1.5см справа")
    print("  ✅ Висячие строки: 28.35pt / -28.35pt")
    print("  ✅ Интервалы: space_before/after 6pt, line_spacing 1.2")
    print("  ✅ Calibri 12pt (НЕ Arial)")
    print("  ✅ Цветные границы #5F497A")
    print("  ✅ Ширины колонок (точные)")
    print("  ✅ Белый текст на фиолетовом фоне (заголовки подписей)")
    print("  ✅ Формат денег: 147 450 (пробел)")
    print("  ✅ Формат часов: 23,00 (запятая)")
