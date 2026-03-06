#!/usr/bin/env python3
"""Scrape ratingruneta.ru web rating and produce XLSX."""

import re, json, time, requests, urllib3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

urllib3.disable_warnings()

HEADERS = {'User-Agent': 'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/120.0.0.0 Safari/537.36'}
RATING_NAME = 'Рейтинг «Разработчики сайтов и веб-сервисов» — 2025'

COST_MAP = {'₽₽₽': 'Высокий', '₽₽': 'Средний', '₽': 'Низкий'}

EMAIL_EXCLUDE = ['example', '.png', '.jpg', '.svg', 'sentry', 'schema', 'google', 'yandex', 'vk.com', 'facebook', 'instagram', 'wixpress', 'w3.org', 'jquery', 'webpack', 'cloudflare', 'github', 'gravatar', 'wordpress', 'bitrix', 'amocrm', 'jivosite', 'calltouch', 'roistat', 'getbutton', 'carrotquest', 'usedesk', 'mindbox', 'retailcrm']

# Pre-known data for first 10
KNOWN = {
    '/agency-aero/': {'site': 'https://aeroidea.ru/', 'phone': '+7 (499) 263-05-91', 'email': 'brief@aeroidea.ru', 'city': 'Москва'},
    '/agency-agima/': {'site': 'http://www.agima.ru', 'phone': '+7 (495) 981-01-85', 'email': 'info@agima.ru', 'city': 'Москва'},
    '/agency-dalee/': {'site': 'http://www.dalee.ru/', 'phone': '+7 (495) 785 94 33', 'email': 'hello@dalee.ru', 'city': 'Москва'},
    '/agency-digital-lab/': {'site': 'https://digital-lab.ru/', 'phone': '+7(495)136-93-07', 'email': 'info@digital-lab.ru', 'city': 'Москва'},
    '/agency-only-com-ru/': {'site': 'https://only.digital/', 'phone': '+7 495 740 9979', 'email': 'hello@only.digital', 'city': 'Новокузнецк'},
    '/agency-intaro/': {'site': 'http://intaro.ru', 'phone': '(495) 268-02-56', 'email': 'mail@intaro.ru', 'city': 'Москва'},
    '/agency-chipsa/': {'site': 'http://chipsa.ru/', 'phone': '+7 (499) 322-86-54', 'email': 'info@chipsa.ru', 'city': 'Красноярск'},
    '/agency-terabit/': {'site': 'https://terabit.ai/', 'phone': '+7 (495) 260-17-06', 'email': 'info@terabit.ai', 'city': 'Москва'},
    '/agency-uplab/': {'site': 'https://www.uplab.ru/', 'phone': '+7 499 653 78 83', 'email': 'info@uplab.ru', 'city': 'Москва'},
    '/agency-purrweb/': {'site': 'https://www.purrweb.com/ru/', 'phone': '+79620392994', 'email': 'hi@purrweb.com', 'city': 'Омск'},
}

companies = []

def parse_companies():
    raw = """Сложные
1 АЭРО ₽₽₽ 72 42 /agency-aero/
2 AGIMA ₽₽₽ 44 27 /agency-agima/
3 Далее/ ₽₽₽ 49 21 /agency-dalee/
4 Digital Lab ₽₽₽ 62 41 /agency-digital-lab/
5 Only ₽₽₽ 38 14 /agency-only-com-ru/
6 Интаро ₽₽₽ 33 21 /agency-intaro/
7 CHIPSA ₽₽₽ 28 2 /agency-chipsa/
8 Terabit Digital ₽₽₽ 24 0 /agency-terabit/
9 Uplab ₽₽ 48 4 /agency-uplab/
10 Purrweb ₽₽₽ 24 0 /agency-purrweb/
11 Escape Tech ₽₽₽ 9 1 /agency-escape-team/
12 «Вебпрактик» ₽₽₽ 24 8 /agency-webpractik/
13 ASTRIO ₽₽₽ 22 5 /agency-astrio/
14 Атвинта ₽₽₽ 39 7 /agency-a42/
15 Qtim ₽₽₽ 13 0 /agency-qtim/
16 NewIT ₽₽ 30 0 /agency-newsite/
17 Creative team ₽₽ 31 4 /agency-ctm/
18 CreativePeople + Piratecode ₽₽₽ 27 8 /agency-cpeople/
19 Riverstart ₽₽ 26 2 /agency-riverstart/
20 Триада ₽₽₽ 9 3 /agency-triadasite/
21 Umbrella IT ₽₽₽ 12 5 /agency-umbrellait/
22 JetStyle ₽₽₽ 11 1 /agency-jetstyle/
23 Heads&Hands ₽₽₽ 12 6 /agency-handh/
24 PROFSOFT ₽₽₽ 15 3 /agency-profsoft/
25 Айтигро ₽₽₽ 10 0 /agency-itgro/
26 NOMADS ₽ 8 3 /agency-nomads/
Средние
1 Космос-Веб ₽₽ 99 9 /agency-cosmos-web/
2 Progressive Media ₽₽ 26 8 /agency-prmedia/
3 Liqium ₽₽ 103 10 /agency-liqium/
4 Цифровой Элемент ₽₽ 66 30 /agency-d-element/
5 Oxem ₽₽ 51 1 /agency-oxem/
6 Braind ₽₽ 32 3 /agency-braind-agency/
7 Red Collar ₽₽ 36 8 /agency-redcollar/
8 Клаудмил ₽₽ 43 4 /agency-cloudmill/
9 its.agency ₽₽ 39 5 /agency-itstudio/
10 AIR Production ₽₽ 63 7 /agency-air-ru-com/
11 MediaTen ₽₽ 31 1 /agency-mediaten/
12 365 Media Group + БЦТ ₽₽ 45 1 /agency-alterwebstudio/
13 Daccel.tech (ex-A5) ₽₽ 33 1 /agency-agency-5/
14 GRCH ₽ 60 7 /agency-grechka-digital/
15 Вебпространство ₽ 48 3 /agency-webprostranstvo/
16 Unistory ₽₽₽ 21 0 /agency-unistory-app/
17 KLBR ₽₽ 23 2 /agency-kolibri-group/
18 aim ₽₽₽ 12 3 /agency-in-aim/
19 SALAMAT ₽₽ 39 0 /agency-websalamat/
20 iMedia Solutions ₽₽₽ 25 1 /agency-imedia-by/
21 BESTWEB ₽ 68 1 /agency-bestweb/
22 Mark Weber ₽₽ 29 1 /agency-markweber/
23 Эмбаси ₽₽ 29 1 /agency-embacy/
24 Pitcher ₽₽ 32 2 /agency-pitcher-agency/
25 INDEXIS ₽₽ 21 4 /agency-indexis/
26 МЭЙК ₽₽ 43 3 /agency-itfnc/
27 Xpage ₽₽ 35 3 /agency-xpage/
28 ВнешнийКод ₽₽ 21 1 /agency-outcode/
29 Текарт ₽₽ 43 9 /agency-web-techart/
30 Киберия ₽₽ 25 0 /agency-cyberia-studio/
31 MACHAON ₽₽ 16 2 /agency-machaon/
32 Студия Т ₽₽ 15 2 /agency-tdsgn/
33 Альт Айти ₽₽ 31 0 /agency-aliterax/
34 KTS ₽₽ 29 0 /agency-kts-tech/
35 SLAM ₽₽ 44 2 /agency-slam-by/
36 АРТВЕЛЛ ₽₽ 20 6 /agency-artwell/
37 FEIP ₽₽ 31 1 /agency-feip/
38 ROSO ₽₽ 46 1 /agency-rosogroup/
39 WEBSHOP ₽₽ 25 0 /agency-webshop-nn/
40 REBOOT DIGITAL STUDIO ₽₽ 11 1 /agency-rebooot-me/
41 Alto ₽₽ 23 0 /agency-alto-codes/
42 ITSPACE.GROUP ₽₽ 12 2 /agency-itspace-group/
43 House ₽ 30 1 /agency-housevl/
44 TRIBE studio ₽₽ 13 2 /agency-studiotribe/
45 Avenue Media ₽₽ 15 0 /agency-avenuemedia/
46 Наумедиа ₽₽ 8 2 /agency-nowmedia/
47 Asanov Agency ₽₽ 14 0 /agency-asanov-agency/
48 NAN agency ₽ 14 1 /agency-nanagency/
49 BrandStudio ₽ 4 4 /agency-bstd/
50 Медиа Лайн ₽ 56 0 /agency-medialine-by/
51 Wemakefab ₽₽ 12 0 /agency-wemakefab/
52 Coalla Agency ₽ 23 0 /agency-coalla/
53 Else ex.ITFactory ₽₽ 21 2 /agency-itfactory/
54 Grokhotov Studio ₽₽ 17 0 /agency-positron-it/
55 «Компот» ₽₽ 23 0 /agency-kompot-bz/
56 PetrogradWeb ₽₽ 23 0 /agency-petrogradweb/
57 BIK.Agency ₽₽₽ 12 1 /agency-bikstudio/
58 ILAR ₽₽ 15 1 /agency-ilartech/
59 Четвёртый Рим ₽₽ 16 1 /agency-4rome/
60 Taptima ₽ 20 1 /agency-taptima/
61 SLT ₽₽ 21 0 /agency-seolt/
62 Аспирити ₽₽₽ 10 1 /agency-aspirity/
63 Пиробайт ₽₽ 15 0 /agency-pyrobyte/
64 ADN ₽₽ 18 1 /agency-a-dn/
65 HAY.agency ₽ 14 0 /agency-hay-studio/
66 OnePix ₽₽ 18 0 /agency-onepix/
67 ARTROCKETS ₽₽ 10 2 /agency-artrockets/
68 Nineseven ₽₽ 11 2 /agency-nineseven/
69 IBRUSH ₽₽₽ 11 2 /agency-ibrush/
70 BeaversBrothers ₽ 12 3 /agency-beaversbrothers/
71 Рецифра ₽ 16 0 /agency-recifra/
72 Fruitful Solutions ₽₽ 14 1 /agency-fru-it/
73 Spectr ₽₽₽ 5 0 /agency-digital-spectr/
74 DNA Team ₽₽ 9 1 /agency-dnateam/
75 OSMI IT ₽₽ 13 0 /agency-myosminozhka/
76 Nomadic Soft ₽₽ 12 0 /agency-nomadicsoft/
77 INSAIM ₽ 11 0 /agency-insaim/
78 OpenStart ₽ 12 1 /agency-openstart/
79 Фьюче ₽₽ 7 1 /agency-future-group/
Простые
1 MEDIA WORKS ₽ 181 20 /agency-mworx/
2 div. ₽ 94 0 /agency-div-pro/
3 АБВ сайт ₽ 132 6 /agency-abcwww/
4 Молния ₽ 128 2 /agency-flashtilda/
5 е—б эдженси ₽ 53 1 /agency-e-b-agency/
6 No Logo Studio ₽ 208 0 /agency-nologostudio/
7 AVA digital ₽ 33 0 /agency-avrora-in/
8 ИСКРА ₽ 28 2 /agency-iskra-st/
9 Lince ₽ 42 4 /agency-ralince/
10 Акцепт ₽ 236 0 /agency-ac-u/
11 ONE PAGE ₽ 31 3 /agency-one-page-site/
12 WebCanape ₽ 173 0 /agency-twinscom/
13 MIKHAILOV studio ₽ 50 0 /agency-mikhailov-studio/
14 NAJES ₽ 41 1 /agency-naje-s/
15 Business-up.org ₽ 64 1 /agency-business-up/
16 Мегагрупп.ру ₽ 356 0 /agency-megagroup/
17 RedKrab ₽ 33 0 /agency-redkrab/
18 Punks ₽ 50 1 /agency-punksdesign/
19 Lucky Site ₽₽ 25 0 /agency-lucky-site/
20 «Вятка IT» ₽ 75 0 /agency-vyatka-it/
21 Vegas ₽ 66 0 /agency-vegas-dev/
22 Восемь Планет ₽ 61 0 /agency-8-planet/
23 EFFECT ₽ 221 0 /agency-effect-su/
24 Komarovaeee ₽ 37 0 /agency-komarovaeee/
25 Bauns ₽ 20 1 /agency-bauns/
26 «Соль» ₽₽ 14 1 /agency-websalt/
27 «ДВИГА» ₽ 30 0 /agency-dviga-marketing/
28 NetLab Creative Studio ₽ 8 0 /agency-netlab-com/
29 INTRID ₽ 30 2 /agency-intrid/
30 Bquadro ₽ 23 2 /agency-bquadro/
31 Belberry ₽ 46 0 /agency-belberry/
32 Migra Software ₽ 10 0 /agency-migra/
33 RHINO DIGITAL ₽ 23 0 /agency-rhino-digital/
34 RUFORMAT® ₽ 19 0 /agency-ruformat/
35 ASMART ₽ 98 1 /agency-asmart-group/
36 STUDIO 512 ₽ 31 0 /agency-ws512/
37 MARS BRANDING ₽ 32 0 /agency-mars-branding/
38 InterCom Digital ₽ 21 0 /agency-icmy-intercom-media/
39 RUSO ₽ 32 0 /agency-rusodot/
40 MRKTNG.bz ₽ 21 0 /agency-mrktng-bz/
41 «Идеалим» ₽ 28 0 /agency-idealim/
42 ipos.digital ₽ 50 0 /agency-1pos/
43 OKTTA ₽ 23 1 /agency-oktta/
44 Кибер-Невод ₽ 23 1 /agency-cyber-nevod/
45 WebMedia ₽ 28 0 /agency-webmedia39/
46 Дессайтс-бай ₽ 76 0 /agency-dessites-by/
47 Lucky-Leads ₽ 22 0 /agency-luckyleads/
48 Электрон ₽ 22 0 /agency-web-electron/
49 «Технологии успеха» ₽ 33 1 /agency-webtu/
50 Storedev/Shopdev Systems ₽ 34 0 /agency-shopdev24/
51 Rizon ₽ 45 1 /agency-rizoncompany/
52 WAYDEV ₽ 44 1 /agency-waydev/
53 AXI ₽ 37 0 /agency-web-axioma/
54 FishCode ₽ 22 0 /agency-fishcode/
55 Russian Robotics ₽₽ 16 1 /agency-rusrobots/
56 INSITE group ₽ 21 0 /agency-insite-group/
57 Art Performance ₽ 15 3 /agency-art-performance/
58 Leto.Website ₽ 13 0 /agency-leto-website/
59 MYOD.IT ₽₽ 5 0 /agency-myod-it/
60 GLADKOV COMPANY ₽ 5 0 /agency-gladkov-company/
61 iT-Wizards ₽ 29 0 /agency-it-wizards/
62 OWL agency ₽ 9 0 /agency-owlagency/
63 500na700 ₽ 17 1 /agency-500na700/
64 Makeit ₽ 10 0 /agency-makeit-da/
65 R.Class ₽₽ 10 0 /agency-rclass/
66 Виктори ₽ 21 1 /agency-victory-su/
67 Delaweb ₽ 10 0 /agency-delaweb/
68 Анкора.Агентство ₽ 12 0 /agency-ancora-agency/
69 Secret Agents ₽ 17 0 /agency-secret-agents/
70 AVOdigital ₽ 17 0 /agency-sozdanie-world/
71 ART6 ₽ 46 0 /agency-art6/
72 Экспресс лаб ₽ 26 1 /agency-expresslab/
73 Kuznets ₽ 12 2 /agency-kuznets-agency/
74 wave ₽ 26 2 /agency-we-wave/
75 Юнона ₽ 28 0 /agency-juno/
76 TRAFF ₽ 24 0 /agency-traff-agency/
77 «Very-good» ₽ 43 0 /agency-very-good/
78 Top-7 ₽ 35 0 /agency-top-7/
79 Nine Squares ₽ 35 0 /agency-ninesquares/
80 АйТи-Эскорт ₽₽ 7 0 /agency-it-escort/
81 Bewave ₽ 13 0 /agency-bewave/
82 We Wizards ₽ 7 0 /agency-wewizards/
83 Vedita ₽ 28 0 /agency-vedita/
84 Scroll Media Solutions ₽ 21 2 /agency-scroll-by/
85 «Импульс» ₽ 38 0 /agency-impulse-marketing/
86 FACE FAMILY ₽₽ 11 1 /agency-facedigital/
87 Flami ₽ 9 0 /agency-flami-agency/
88 Старта ₽ 19 0 /agency-starta/
89 Мир сайтов ₽ 22 2 /agency-mirsaitov/
90 German Web ₽ 18 0 /agency-german-web/
91 «Инфо-Сити» ₽ 20 0 /agency-info-city/
92 «АртФактор» ₽ 62 0 /agency-artfactor/
93 LYZLOV DIGITAL MARKETING ₽ 16 0 /agency-lyzlovdm/
94 Scrum studio White ₽ 8 0 /agency-scrum-studio/
95 БОБРЫ ₽ 63 0 /agency-bobers/
96 ENJOY TOUCH ₽ 11 0 /agency-enjoytouch/
97 Golden Studio ₽ 19 0 /agency-golden-studio/
98 Импульс-Маркетинг ₽ 13 0 /agency-impulsit/
99 FACEDIGITAL ₽ 25 1 /agency-facedigital-agency/
100 TopForm ₽₽ 8 1 /agency-topform/
101 «Простые решения» ₽ 27 1 /agency-ima-pr/
102 Х10 Agency ₽ 25 0 /agency-x10/
103 Stark Consulting Agency ₽ 29 0 /agency-stark-team/
104 Astra Marketing ₽ 27 0 /agency-astra-agency/
105 Seo Move ₽ 35 0 /agency-seo-move/
106 Media Bay ₽ 16 0 /agency-media-bay/
107 ТЛТ ПРО ₽ 11 0 /agency-nethouse-altrp/
108 GM Group ₽ 35 0 /agency-gmorning/
109 TOLA ₽ 16 0 /agency-tolacompany/
110 Pro Pixeli ₽ 7 0 /agency-propixeli-studio/
111 «Всё отлично!» ₽ 17 0 /agency-aplayweb/
112 РУВС ₽ 11 0 /agency-russiaws/
113 «Сайт PRO» ₽ 22 0 /agency-sitepro-pro/
114 L-Digital ₽ 14 0 /agency-l-digital/
115 DATAKIT ₽ 18 0 /agency-datakit/
116 Clockwork ₽ 8 0 /agency-clockwork/
117 MediaMint ₽ 12 0 /agency-mediamint/
118 Рикорда ₽ 22 0 /agency-rikorda-group/
119 АЙТИМОД ₽ 10 0 /agency-itmod/
120 Артметрика ₽ 12 0 /agency-arkhipov-digital/
121 Винтра ₽ 21 0 /agency-wintramedia/
122 SO.USE ₽ 15 0 /agency-so-use/
123 Craft Promotion ₽ 11 0 /agency-craftpromotion/
124 BLOT.PRO ₽ 8 0 /agency-blot-pro/
125 Denero ₽ 13 0 /agency-denero/
126 Klatcen ₽ 14 0 /agency-klatcenklatcen/
127 KOTTE ₽ 23 0 /agency-kotte-agency/
128 ASAP ₽ 13 0 /agency-asap-ag/
129 Port Digital ₽ 8 0 /agency-snw-digital/
130 CAPYBARA digital ₽ 13 0 /agency-capyba/
131 Сайт-Креатив ₽ 17 0 /agency-site-creative/
132 Praweb ₽ 15 0 /agency-praweb/
133 Comly ₽ 13 0 /agency-comly-dev/
134 REDLINE ₽ 9 0 /agency-redline-by/
135 Zavod ₽ 15 0 /agency-zavod-team/
136 Барбарис ₽ 13 0 /agency-dsbarberry/
137 ZLS ₽ 11 0 /agency-zls/
138 Baitex ₽ 5 0 /agency-baitex/
139 Linkodium ₽ 12 0 /agency-linkodium/
140 Legend ₽ 16 0 /agency-1leg/
141 Daimax Digital ₽ 16 0 /agency-popov-promo/
142 A&A ₽ 9 0 /agency-it-uu/
143 Anycase-IT ₽ 16 0 /agency-anykey-it/
144 Пластилин-арт ₽ 12 0 /agency-plastilin-art/
145 МЕТАШАРКС ₽ 10 1 /agency-metasharks/
146 CULT ₽ 17 0 /agency-cult-digital/
147 ИСТИНА В МАРКЕТИНГЕ ₽ 13 0 /agency-istina-marketinga-rf/
148 360i ₽ 13 0 /agency-360i/
149 Волшебный Ёжик ₽ 12 0 /agency-vezhik-by/
150 Lead Wave ₽ 7 0 /agency-lead-wave/
151 Webnavigator ₽ 7 0 /agency-webnavigator/
152 Mobility.top ₽ 9 0 /agency-mobility-top/"""

    category = None
    for line in raw.strip().split('\n'):
        line = line.strip()
        if line in ('Сложные', 'Средние', 'Простые'):
            category = line
            continue
        # Parse: rank name cost clients big_clients slug
        m = re.match(r'^(\d+)\s+(.+?)\s+(₽{1,3})\s+(\d+)\s+(\d+)\s+(/agency-[^/]+/)\s*$', line)
        if m:
            companies.append({
                'category': category,
                'rank': int(m.group(1)),
                'name': m.group(2),
                'cost_raw': m.group(3),
                'cost': COST_MAP[m.group(3)],
                'clients': int(m.group(4)),
                'big_clients': int(m.group(5)),
                'slug': m.group(6),
            })

    print(f"Parsed {len(companies)} companies")

def fetch_contact_info(slug):
    """Get phone, city, site from ratingruneta contacts page."""
    if slug in KNOWN:
        return KNOWN[slug]
    
    url = f'https://ratingruneta.ru{slug}contacts/'
    try:
        r = requests.get(url, headers=HEADERS, timeout=15, verify=False)
        r.raise_for_status()
        html = r.text
        idx = html.find('window.__INITIAL_STATE__')
        if idx == -1:
            return {}
        start = html.index('{', idx)
        # Find matching closing brace
        depth = 0
        for i in range(start, len(html)):
            if html[i] == '{': depth += 1
            elif html[i] == '}': depth -= 1
            if depth == 0:
                json_str_raw = html[start:i+1]
                break
        else:
            return {}
        json_str = json_str_raw
        # Fix invalid JSON escapes like \< \> \& etc
        json_str = re.sub(r'\\(?!["\\/bfnrtu])', r'\\\\', json_str)
        
        try:
            data = json.loads(json_str)
        except json.JSONDecodeError as e:
            print(f"  JSON parse error: {e}")
            return {}
        
        header = data.get('agency', {}).get('header', {}).get('data', {})
        result = {}
        
        label = header.get('label', {})
        if label.get('link'):
            result['site'] = label['link']
        
        info = header.get('info', {})
        if info.get('phone'):
            result['phone'] = info['phone']
        
        cities = info.get('cities', [])
        if cities and cities[0].get('city'):
            result['city'] = cities[0]['city']
        
        return result
    except Exception as e:
        print(f"  Error fetching {url}: {e}")
        return {}

def find_email(site_url):
    """Try to find email on company website."""
    if not site_url:
        return ''
    
    # Normalize
    if not site_url.startswith('http'):
        site_url = 'https://' + site_url
    
    urls_to_try = [site_url]
    base = site_url.rstrip('/')
    urls_to_try.append(base + '/contacts')
    urls_to_try.append(base + '/contact')
    urls_to_try.append(base + '/contacts/')
    urls_to_try.append(base + '/about')
    
    for url in urls_to_try:
        try:
            r = requests.get(url, headers=HEADERS, timeout=10, verify=False, allow_redirects=True)
            if r.status_code != 200:
                continue
            text = r.text
            emails = re.findall(r'[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}', text)
            for email in emails:
                email_lower = email.lower()
                if any(exc in email_lower for exc in EMAIL_EXCLUDE):
                    continue
                if email_lower.endswith(('.png', '.jpg', '.svg', '.gif', '.webp', '.css', '.js')):
                    continue
                return email
        except:
            continue
    return ''

def generate_xlsx(companies, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Рейтинг 2025'
    
    headers = [
        'Название компании', 'Название рейтинга', 'Место в рейтинге',
        'Ссылка на сайт', 'Страница в рейтинге', 'Телефон', 'Email',
        'Город', 'Категория', 'Уровень стоимости', 'Кол-во заказчиков',
        'Кол-во крупных заказчиков'
    ]
    
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    alt_fill = PatternFill(start_color='DEEAF1', end_color='DEEAF1', fill_type='solid')
    link_font = Font(color='0563C1', underline='single')
    thin_border = Border(
        left=Side(style='thin', color='B4C6E7'),
        right=Side(style='thin', color='B4C6E7'),
        top=Side(style='thin', color='B4C6E7'),
        bottom=Side(style='thin', color='B4C6E7')
    )
    
    # Headers
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    
    ws.freeze_panes = 'A2'
    
    for i, c in enumerate(companies):
        row = i + 2
        rating_url = f"https://ratingruneta.ru{c['slug']}"
        
        ws.cell(row=row, column=1, value=c['name']).border = thin_border
        ws.cell(row=row, column=2, value=RATING_NAME).border = thin_border
        ws.cell(row=row, column=3, value=c['rank']).border = thin_border
        
        # Site link
        site = c.get('site', '')
        cell4 = ws.cell(row=row, column=4, value=site)
        cell4.border = thin_border
        if site:
            cell4.hyperlink = site
            cell4.font = link_font
        
        # Rating page link
        cell5 = ws.cell(row=row, column=5, value=rating_url)
        cell5.hyperlink = rating_url
        cell5.font = link_font
        cell5.border = thin_border
        
        ws.cell(row=row, column=6, value=c.get('phone', '')).border = thin_border
        ws.cell(row=row, column=7, value=c.get('email', '')).border = thin_border
        ws.cell(row=row, column=8, value=c.get('city', '')).border = thin_border
        ws.cell(row=row, column=9, value=c['category']).border = thin_border
        ws.cell(row=row, column=10, value=c['cost']).border = thin_border
        ws.cell(row=row, column=11, value=c['clients']).border = thin_border
        ws.cell(row=row, column=12, value=c['big_clients']).border = thin_border
        
        # Alternating row colors
        if row % 2 == 0:
            for col in range(1, 13):
                ws.cell(row=row, column=col).fill = alt_fill
    
    # Column widths
    widths = [30, 45, 8, 35, 40, 22, 28, 18, 12, 15, 12, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    
    ws.row_dimensions[1].height = 30
    
    wb.save(output_path)
    print(f"Saved {output_path}")

def main():
    parse_companies()
    
    total = len(companies)
    email_count = 0
    
    for i, c in enumerate(companies):
        slug = c['slug']
        print(f"[{i+1}/{total}] {c['name']} ({slug})")
        
        info = fetch_contact_info(slug)
        c['site'] = info.get('site', '')
        c['phone'] = info.get('phone', '')
        c['city'] = info.get('city', '')
        
        if slug in KNOWN:
            c['email'] = KNOWN[slug].get('email', '')
        else:
            c['email'] = find_email(c['site'])
            time.sleep(0.3)
        
        if c['email']:
            email_count += 1
        
        if (i + 1) % 20 == 0:
            print(f"  Progress: {i+1}/{total}, emails found: {email_count}")
        
        time.sleep(0.2)
    
    output = '/home/clawdbot/.openclaw/workspace/ratingruneta_full.xlsx'
    generate_xlsx(companies, output)
    
    print(f"\n=== DONE ===")
    print(f"Total companies: {total}")
    print(f"Emails found: {email_count}")
    print(f"File: {output}")

if __name__ == '__main__':
    main()
