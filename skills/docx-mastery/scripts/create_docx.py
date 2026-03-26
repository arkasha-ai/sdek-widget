#!/usr/bin/env python3
"""Quick DOCX creator from JSON spec.

Usage:
    python create_docx.py spec.json output.docx
    
Spec format:
{
    "title": "Document Title",
    "font": "Times New Roman",
    "font_size": 12,
    "margins": {"left": 3, "right": 1.5, "top": 2, "bottom": 2},
    "content": [
        {"type": "heading", "level": 1, "text": "Chapter 1"},
        {"type": "paragraph", "text": "Body text here", "bold": false},
        {"type": "table", "headers": ["Col1", "Col2"], "rows": [["a", "b"]]},
        {"type": "image", "path": "img.png", "width_cm": 10},
        {"type": "page_break"},
        {"type": "list", "style": "bullet", "items": ["Item 1", "Item 2"]}
    ]
}
"""

import json
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_document(spec: dict, output_path: str):
    doc = Document()
    
    # Page setup
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    margins = spec.get("margins", {"left": 2, "right": 1.5, "top": 2, "bottom": 2})
    section.left_margin = Cm(margins.get("left", 2))
    section.right_margin = Cm(margins.get("right", 1.5))
    section.top_margin = Cm(margins.get("top", 2))
    section.bottom_margin = Cm(margins.get("bottom", 2))
    
    # Default font
    style = doc.styles['Normal']
    style.font.name = spec.get("font", "Times New Roman")
    style.font.size = Pt(spec.get("font_size", 12))
    
    # Title
    if "title" in spec:
        p = doc.add_heading(spec["title"], level=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Content
    for item in spec.get("content", []):
        t = item.get("type", "paragraph")
        
        if t == "heading":
            doc.add_heading(item["text"], level=item.get("level", 1))
        
        elif t == "paragraph":
            p = doc.add_paragraph()
            run = p.add_run(item["text"])
            if item.get("bold"):
                run.bold = True
            if item.get("italic"):
                run.italic = True
            if item.get("align") == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif item.get("align") == "right":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        elif t == "table":
            headers = item.get("headers", [])
            rows = item.get("rows", [])
            table = doc.add_table(rows=1 + len(rows), cols=len(headers), style='Table Grid')
            # Headers
            for i, h in enumerate(headers):
                cell = table.cell(0, i)
                cell.text = h
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            # Data
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row):
                    table.cell(ri + 1, ci).text = str(val)
        
        elif t == "image":
            width = Cm(item.get("width_cm", 10)) if "width_cm" in item else None
            doc.add_picture(item["path"], width=width)
        
        elif t == "page_break":
            doc.add_page_break()
        
        elif t == "list":
            style_name = "List Bullet" if item.get("style") == "bullet" else "List Number"
            for li in item.get("items", []):
                doc.add_paragraph(li, style=style_name)
    
    doc.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python create_docx.py spec.json output.docx")
        sys.exit(1)
    
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    create_document(spec, sys.argv[2])
