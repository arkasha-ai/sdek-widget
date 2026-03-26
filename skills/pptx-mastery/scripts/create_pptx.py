#!/usr/bin/env python3
"""Quick PPTX creator from JSON spec.

Usage:
    python create_pptx.py spec.json output.pptx

Spec format:
{
    "title": "Presentation Title",
    "subtitle": "Company Name",
    "theme": "blue",
    "slides": [
        {"type": "title", "title": "Main Title", "subtitle": "Subtitle"},
        {"type": "content", "title": "Slide Title", "bullets": ["Point 1", "Point 2"]},
        {"type": "image", "title": "Screenshot", "image": "path.png"},
        {"type": "table", "title": "Data", "headers": ["A","B"], "rows": [["1","2"]]},
        {"type": "blank"}
    ]
}
"""

import json
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

THEMES = {
    "blue": {"primary": RGBColor(0x44, 0x72, 0xC4), "text": RGBColor(0x33, 0x33, 0x33)},
    "dark": {"primary": RGBColor(0x1A, 0x1A, 0x2E), "text": RGBColor(0xFF, 0xFF, 0xFF)},
    "gravity": {"primary": RGBColor(0x2C, 0x3E, 0x50), "text": RGBColor(0x2C, 0x3E, 0x50)},
}


def create_presentation(spec: dict, output_path: str):
    prs = Presentation()
    prs.slide_width = Cm(33.867)
    prs.slide_height = Cm(19.05)
    
    theme = THEMES.get(spec.get("theme", "blue"), THEMES["blue"])
    
    for slide_spec in spec.get("slides", []):
        stype = slide_spec.get("type", "content")
        
        if stype == "title":
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = slide_spec.get("title", "")
            if slide.placeholders.get(1):
                slide.placeholders[1].text = slide_spec.get("subtitle", "")
        
        elif stype == "content":
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = slide_spec.get("title", "")
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.clear()
            for i, bullet in enumerate(slide_spec.get("bullets", [])):
                if i == 0:
                    tf.paragraphs[0].text = bullet
                else:
                    p = tf.add_paragraph()
                    p.text = bullet
                    p.level = slide_spec.get("level", 0)
        
        elif stype == "image":
            slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
            slide.shapes.title.text = slide_spec.get("title", "")
            img_path = slide_spec.get("image", "")
            if img_path and Path(img_path).exists():
                slide.shapes.add_picture(img_path, Inches(1), Inches(2), width=Inches(8))
        
        elif stype == "table":
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = slide_spec.get("title", "")
            headers = slide_spec.get("headers", [])
            rows = slide_spec.get("rows", [])
            if headers:
                tbl = slide.shapes.add_table(
                    1 + len(rows), len(headers),
                    Inches(0.5), Inches(2), Inches(9), Inches(0.5 + 0.4 * len(rows))
                ).table
                for ci, h in enumerate(headers):
                    tbl.cell(0, ci).text = h
                for ri, row in enumerate(rows):
                    for ci, val in enumerate(row):
                        tbl.cell(ri + 1, ci).text = str(val)
        
        elif stype == "blank":
            prs.slides.add_slide(prs.slide_layouts[6])
        
        # Speaker notes
        if "notes" in slide_spec:
            slide.notes_slide.notes_text_frame.text = slide_spec["notes"]
    
    prs.save(output_path)
    print(f"Saved: {output_path} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python create_pptx.py spec.json output.pptx")
        sys.exit(1)
    
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    create_presentation(spec, sys.argv[2])
