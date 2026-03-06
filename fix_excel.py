#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import openpyxl
from openpyxl import load_workbook

def fix_formulas(filename):
    print("=== АНАЛИЗ И ИСПРАВЛЕНИЕ ФОРМУЛ ===")
    
    # Загружаем workbook без вычисления формул, чтобы увидеть исходные данные
    wb = load_workbook(filename, data_only=False)
    
    print("\n=== ПРОВЕРКА ИСХОДНЫХ ДАННЫХ ===")
    ws_dec = wb['Декомпозиция']
    ws_spec = wb['Специалисты']
    
    # Получаем ставки специалистов
    rates = {}
    for row in range(4, 10):  # строки 4-9
        role = ws_spec[f'A{row}'].value
        rate = ws_spec[f'B{row}'].value
        if role and rate:
            rates[role] = rate
    print(f"Ставки специалистов: {rates}")
    
    # Проверяем первые несколько строк с данными
    print("\nПроверка исходных данных (часы по специалистам):")
    for row in range(4, 12):  # первые 8 строк
        task_num = ws_dec[f'A{row}'].value
        task = ws_dec[f'C{row}'].value
        
        # Проверяем часы по специалистам (столбцы E-J)
        hours_data = {}
        for col_letter, col_idx in [('E', 5), ('F', 6), ('G', 7), ('H', 8), ('I', 9), ('J', 10)]:
            hours = ws_dec[f'{col_letter}{row}'].value
            hours_data[col_letter] = hours
        
        # Формулы
        sum_formula = ws_dec[f'K{row}'].value
        cost_formula = ws_dec[f'L{row}'].value
        
        print(f"\nСтрока {row}: {task_num} - {task}")
        print(f"  Часы по специалистам: {hours_data}")
        print(f"  Формула суммы (K): {sum_formula}")
        print(f"  Формула стоимости (L): {cost_formula}")
    
    # Попробуем пересчитать вручную для первых строк
    print("\n=== РУЧНОЙ ПЕРЕСЧЕТ ===")
    for row in range(4, 8):
        task_num = ws_dec[f'A{row}'].value
        task = ws_dec[f'C{row}'].value
        
        # Получаем часы
        e_hours = ws_dec[f'E{row}'].value or 0
        f_hours = ws_dec[f'F{row}'].value or 0  
        g_hours = ws_dec[f'G{row}'].value or 0
        h_hours = ws_dec[f'H{row}'].value or 0
        i_hours = ws_dec[f'I{row}'].value or 0
        j_hours = ws_dec[f'J{row}'].value or 0
        
        # Пересчитываем вручную
        total_hours = e_hours + f_hours + g_hours + h_hours + i_hours + j_hours
        
        # Стоимость по ставкам
        total_cost = (e_hours * rates.get('ML-инженер', 0) + 
                     f_hours * rates.get('Backend Senior', 0) +
                     g_hours * rates.get('Frontend', 0) +
                     h_hours * rates.get('DevOps', 0) +
                     i_hours * rates.get('QA', 0) +
                     j_hours * rates.get('РП', 0))
        
        print(f"\nСтрока {row}: {task_num}")
        print(f"  Исходные часы: E={e_hours}, F={f_hours}, G={g_hours}, H={h_hours}, I={i_hours}, J={j_hours}")
        print(f"  Итого часов: {total_hours}")
        print(f"  Итого стоимость: {total_cost:,} ₽")
        
        # Проверяем формулы на корректность
        expected_sum_formula = f"=SUM(E{row}:J{row})"
        expected_cost_formula = (f"=E{row}*Специалисты!$B$4+F{row}*Специалисты!$B$5+G{row}*Специалисты!$B$6+"
                               f"H{row}*Специалисты!$B$7+I{row}*Специалисты!$B$8+J{row}*Специалисты!$B$9")
        
        current_sum_formula = ws_dec[f'K{row}'].value
        current_cost_formula = ws_dec[f'L{row}'].value
        
        if current_sum_formula != expected_sum_formula:
            print(f"  ИСПРАВЛЯЮ формулу суммы: {current_sum_formula} -> {expected_sum_formula}")
            ws_dec[f'K{row}'] = expected_sum_formula
        else:
            print(f"  Формула суммы корректна")
            
        if current_cost_formula != expected_cost_formula:
            print(f"  ИСПРАВЛЯЮ формулу стоимости")
            ws_dec[f'L{row}'] = expected_cost_formula
        else:
            print(f"  Формула стоимости корректна")
    
    # Сохраняем исправленный файл
    output_filename = filename.replace('.xlsx', '_fixed.xlsx')
    wb.save(output_filename)
    wb.close()
    
    print(f"\nСохранен исправленный файл: {output_filename}")
    
    return output_filename

def recalculate_all_rows(filename):
    print("\n=== ПЕРЕСЧЕТ ВСЕХ СТРОК ===")
    
    wb = load_workbook(filename, data_only=False)
    ws_dec = wb['Декомпозиция']
    ws_spec = wb['Специалисты']
    
    # Получаем ставки
    rates = {}
    for row in range(4, 10):
        role = ws_spec[f'A{row}'].value
        rate = ws_spec[f'B{row}'].value
        if role and rate:
            rates[role] = rate
    print(f"Ставки: {rates}")
    
    total_hours = 0
    total_cost = 0
    
    # Пересчитываем все строки с данными
    for row in range(4, 35):  # строки 4-34
        task_num = ws_dec[f'A{row}'].value
        if not task_num:  # пропускаем пустые строки
            continue
            
        # Получаем часы
        e_hours = ws_dec[f'E{row}'].value or 0
        f_hours = ws_dec[f'F{row}'].value or 0
        g_hours = ws_dec[f'G{row}'].value or 0
        h_hours = ws_dec[f'H{row}'].value or 0
        i_hours = ws_dec[f'I{row}'].value or 0
        j_hours = ws_dec[f'J{row}'].value or 0
        
        row_total_hours = e_hours + f_hours + g_hours + h_hours + i_hours + j_hours
        row_total_cost = (e_hours * rates.get('ML-инженер', 0) + 
                         f_hours * rates.get('Backend Senior', 0) +
                         g_hours * rates.get('Frontend', 0) +
                         h_hours * rates.get('DevOps', 0) +
                         i_hours * rates.get('QA', 0) +
                         j_hours * rates.get('РП', 0))
        
        print(f"Строка {row}: {task_num} - {row_total_hours} ч, {row_total_cost:,} ₽")
        
        total_hours += row_total_hours
        total_cost += row_total_cost
        
        # Вставляем рассчитанные значения
        ws_dec[f'K{row}'] = row_total_hours
        ws_dec[f'L{row}'] = row_total_cost
    
    print(f"\nИТОГО: {total_hours} часов, {total_cost:,} ₽")
    
    # Сохраняем файл с пересчитанными значениями
    output_filename = filename.replace('.xlsx', '_calculated.xlsx')
    wb.save(output_filename)
    wb.close()
    
    return output_filename, total_hours, total_cost

if __name__ == "__main__":
    filename = "/home/clawdbot/.openclaw/workspace/drafts/SyncVoice_Decomposition_v1.xlsx"
    
    # Исправляем формулы
    fixed_file = fix_formulas(filename)
    
    # Пересчитываем все значения
    calc_file, total_h, total_c = recalculate_all_rows(fixed_file)
    
    print(f"\nФинальный результат: {total_h} часов, {total_c:,} ₽")
    print(f"Ожидалось: 994 часов, 2,624,000 ₽")
    print(f"Разница: {total_h - 994} часов, {total_c - 2624000:,} ₽")